"""Spreadsheet import and export (CSV, XLSX, ODS).

Import flow (API: ``POST /import/preview`` then ``POST /import/run``):

1. :func:`store_upload` keeps the uploaded file under ``<local data>/import/<token>.<ext>`` (24 h).
2. :func:`preview` reads it (CSV: encoding and delimiter detection incl. German Excel
   ``;``/cp1252; XLSX via openpyxl; ODS via odfpy), detects the header row and suggests
   a column -> field mapping (rapidfuzz over field names, German labels and synonyms).
   Units in the header (``Kapazität (pF)``, ``Widerstand [kOhm]``) become the mapping's
   ``unit_hint``: plain numbers of that column are read in this unit (``100`` -> ``100pF``).
   Percent-formatted cells (Excel/LibreOffice store ``1 %`` as ``0.01``) are read as ``1%``.
3. :func:`run_import` (a background job, kind ``import``) backs the database up and writes the rows with
   :func:`partforge.db.parts.upsert_parts`; its result carries ``error_count``, ``outcome``
   (``ok`` = no rejected row, ``partial`` = rejected rows but parts were created/updated,
   ``failed`` = rejected rows and nothing created/updated) and the final ``message``
   („Import fertig: 120 neu, 4 aktualisiert, 2 abgelehnt“). ``strict`` turns unreadable values
   (no number, wrong unit, no date …) into errors of their row. With ``library`` (API) the job holds
   that library for its whole run (:func:`partforge.jobs.library_hold`): a library switch cancels it,
   and a cancelled import imported nothing (one transaction).

Per-row status: :func:`check_rows` (also part of :func:`preview` when a category is chosen) runs the
import as a dry run (:func:`partforge.db.parts.plan_upsert`) for the current mapping, mode and fixed
values: ``create``/``update``/``unchanged``/``skip``/``reject`` with German reasons per row.

Export: :func:`export_parts` writes a category (all or selected parts, chosen columns)
as CSV (UTF-8 with BOM, ``;`` – opens correctly in German Excel) or XLSX (bold header,
autofilter, frozen header row, sensible column widths). XLSX cells are always literal
text – ``=Value`` (Altium indirection) never becomes a formula. CSV has no cell types:
Excel evaluates fields starting with ``= + - @`` (and tab/CR) when it opens a CSV, so by default
(``excel_safe=True``) such values get an apostrophe in front (the importer removes it again) – values
from datasheets, foreign spreadsheets or other users must never run as formulas. A minus before a
plain measurement stays readable (``-55°C``, ``-40 … 85°C``, ``-100 … +350ppm/K``); ``excel_safe=False``
writes every value unchanged. XLSX values longer than Excel's cell limit (32,767 characters) are cut
by Excel's format; the export reports them (``ExportFile.warnings``, job message) – CSV keeps them. The export job
(:func:`export_to_file` with ``library``) only reads and holds nothing; a library switch ends it
cleanly (``LibraryChanged`` → ``cancelled``) – never an export of, or into, another library.
"""
from __future__ import annotations

import codecs
import csv
import datetime as _dt
import io
import logging
import os
import re
import secrets
import time
from collections import Counter
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field
from rapidfuzz import fuzz, utils

from ..config import local_data_dir
from ..db import backup, categories, parts
from ..db.models import CategoryDef, FieldDef, PartFilter, SortSpec
from ..db.values import PART_NUMBER, parse_field_quantity, to_text, units_compatible
from ..domain.text import plural
from ..jobs import JobCancelled, JobContext, LibraryChanged, cancelled_by_switch, library_hold, open_library

log = logging.getLogger(__name__)

__all__ = [
    "ExportFile",
    "ImportOptions",
    "MappingItem",
    "SheetData",
    "SpreadsheetError",
    "apply_unit_hint",
    "build_rows",
    "check_rows",
    "export_parts",
    "export_to_file",
    "header_unit",
    "import_message",
    "import_summary",
    "preview",
    "read_table_file",
    "render_export",
    "run_import",
    "store_upload",
    "suggest_mapping",
    "upload_path",
]

SUPPORTED_EXTENSIONS = (".csv", ".txt", ".tsv", ".xlsx", ".xlsm", ".ods")
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
UPLOAD_MAX_AGE_S = 24 * 3600
MATCH_THRESHOLD = 78.0
_MAX_ROWS = 200_000
_MAX_COLUMNS = 512
#: Characters Excel keeps in one cell (XLSX): longer texts are cut by the format (openpyxl writes exactly
#: this many) – the export reports them, the import preview warns about cells of exactly this length.
XLSX_CELL_LIMIT = 32_767
_CELL_LIMIT_TEXT = "32.767"  # German spelling for messages
_CSV_FIELD_LIMIT = 2**31 - 1  # Python's csv module refuses fields over 128 KiB by default
_CHECK_EVERY = 2000  # rows between cancellation checks while the import job reads its file
_PROGRESS_EVERY = 500  # rows between progress reports while an export is written
_DELIMITERS = (";", "\t", ",", "|")
_TOKEN_RE = re.compile(r"^[a-f0-9]{24}$")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")  # Excel evaluates CSV fields starting with these
_EXCEL_GUARD = "'"
_GUARDED = re.compile(r"^'+[=+\-@\t\r]")  # guard apostrophe(s) before a formula start
# a negative measurement that Excel shows as text, never a formula: "-55°C", "-40 … 85°C", "-100 … +350ppm/K"
_SAFE_NEGATIVE = re.compile(
    r"^-\d[\d.,]*[^\s=+\-@|!();\"'\\]*"
    r"(?:\s*(?:…|\.\.\.|~|to|bis)\s*[-+±]?\d[\d.,]*[^\s=+\-@|!();\"'\\]*)?$")
# "(pF)", "[kOhm]", "(mW)" in a column header (no digits: "(@100MHz)" is not a unit)
_HEADER_UNIT = re.compile(r"[\(\[]\s*([^\(\)\[\]\d]{1,12}?)\s*[\)\]]")
_PLAIN_NUMBER = re.compile(r"^[-+]?\d+(?:[.,]\d+)?(?:[eE][-+]?\d+)?$")
_RANGE_PARTS = re.compile(r"(\s*(?:\.\.\.|…|\bto\b|\bbis\b|~|\s-\s)\s*)")
# Excel number format: "%" outside quotes/escapes multiplies by 100
_FORMAT_LITERALS = re.compile(r'"[^"]*"|\\.')


class SpreadsheetError(ValueError):
    """Import/export problem with a German message (API: 400)."""


# --------------------------------------------------------------------------- synonyms

# Additional spellings (German, English, distributor exports) per column name.
SYNONYMS: dict[str, list[str]] = {
    "Part Number": ["teilenummer", "artikelnummer", "bauteilnummer", "part no", "part #", "pn", "interne nummer",
                    "item id", "design item id", "lagernummer"],
    "Manufacturer": ["hersteller", "herstellername", "mfr", "mfg", "manufacturer name", "brand", "marke"],
    "Manufacturer Part Number": ["mpn", "hersteller teilenummer", "herstellerteilenummer", "hersteller nr",
                                 "herstellernummer", "hersteller artikelnummer", "mfr part number", "mfr part",
                                 "mfr part no", "mfg part number", "manufacturer part no", "manufacturer part",
                                 "typ", "typenbezeichnung"],
    "Value": ["wert"],
    "Description": ["beschreibung", "bezeichnung", "desc", "info", "artikelbezeichnung", "component description"],
    "Comment": ["kommentar"],
    "Package": ["gehäuse", "gehaeuse", "bauform", "case", "case package", "package case", "device package"],
    "Mounting": ["montage", "montageart", "mounting type", "mount", "bestückungsart"],
    "Height": ["höhe", "hoehe", "bauhöhe", "height max", "height seated max"],
    "Operating Temperature": ["betriebstemperatur", "temperaturbereich", "temp range", "operating temperature range"],
    "Library Ref": ["symbol", "symbolname", "schaltplansymbol", "schematic symbol", "lib ref"],
    "Library Path": ["symbolbibliothek", "schlib", "symbol library"],
    "Footprint Ref": ["footprint", "footprint name", "pcb footprint", "land pattern", "gehäuse footprint"],
    "Footprint Path": ["footprint bibliothek", "pcblib", "footprint library"],
    "ComponentLink1URL": ["datenblatt", "datasheet", "datasheet url", "datenblatt url", "datenblatt link",
                          "datasheet link", "url", "link"],
    "ComponentLink1Description": ["link beschriftung", "link description"],
    "Supplier 1": ["lieferant", "distributor", "supplier", "händler"],
    "Supplier Part Number 1": ["bestellnummer", "bestell nr", "lieferantennummer", "supplier part number",
                               "distributor part number", "order number", "order code", "digikey part number",
                               "digi key part number", "mouser part number", "mouser no", "farnell order code"],
    "Lifecycle": ["lebenszyklus", "lifecycle status", "part status", "produktstatus"],
    "RoHS": ["rohs konform", "rohs compliant", "rohs status"],
    "Notes": ["notizen", "notiz", "bemerkung", "bemerkungen", "anmerkung", "anmerkungen"],
    "Resistance": ["widerstand", "widerstandswert", "ohm"],
    "Capacitance": ["kapazität", "kapazitaet"],
    "Inductance": ["induktivität", "induktivitaet"],
    "Impedance": ["impedanz", "impedance at 100mhz", "impedance 100mhz"],
    "Tolerance": ["toleranz", "tol"],
    "Power": ["leistung", "power rating", "nennleistung", "verlustleistung", "watt"],
    "Voltage": ["spannung", "nennspannung", "voltage rating", "voltage rated", "rated voltage", "voltage dc"],
    "TCR": ["tk", "temperaturkoeffizient", "temperature coefficient", "tempco"],
    "Dielectric": ["dielektrikum", "temperature characteristic", "temperaturcharakteristik"],
    "Size": ["baugröße", "baugroesse", "größe", "case code", "case code imperial", "imperial size"],
    "Current Rating": ["nennstrom", "rated current", "current"],
    "Frequency": ["frequenz"],
    "Color": ["farbe", "colour"],
    "Pins": ["pole", "polzahl", "anzahl pins", "number of pins", "pin count", "positions"],
    "Pitch": ["raster", "rastermaß"],
}


# --------------------------------------------------------------------------- reading


@dataclass
class SheetData:
    """A table read from a file (header row detection applied)."""

    sheets: list[str]
    sheet: str | None
    headers: list[str]
    rows: list[list[Any]]  # data rows below the header (raw cell values)
    header_row: int | None  # 1-based row number of the header, None = no header row
    encoding: str | None = None
    delimiter: str | None = None
    raw_first_rows: list[list[Any]] = field(default_factory=list)
    lines: list[int] = field(default_factory=list)  # 1-based sheet row of every data row (empty rows skipped)


def _decode(data: bytes) -> tuple[str, str]:
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8):].decode("utf-8", errors="replace"), "utf-8-sig"
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace"), "utf-16"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1252"), "cp1252"
    except UnicodeDecodeError:
        return data.decode("latin-1"), "latin-1"


def _sniff_delimiter(text: str) -> str:
    """Delimiter giving the most consistent field count (>1) over the first lines."""
    lines = [line for line in text.splitlines() if line.strip()][:50]
    if not lines:
        return ";"
    best_key: tuple[float, int, int] | None = None
    best_delim = ";"
    for priority, delim in enumerate(_DELIMITERS):
        counts = [len(row) for row in csv.reader(lines, delimiter=delim)]
        if not counts:
            continue
        mode, freq = Counter(counts).most_common(1)[0]
        if mode < 2:
            continue
        # consistency of the field count first, then the delimiter priority (; before ,
        # because German CSVs use the decimal comma), then the number of columns
        key = (freq / len(counts), -priority, mode)
        if best_key is None or key > best_key:
            best_key, best_delim = key, delim
    return best_delim


def _unguard(value: str) -> str:
    """Remove the apostrophe an Excel-safe CSV export put before ``= + - @`` (tab, CR) – exactly one of
    a run of apostrophes before such a character (the export adds one to a value that itself starts with
    ``'=`` …, see :func:`_csv_cell`), so the round trip is lossless."""
    if _GUARDED.match(value):
        return value[1:]
    return value


def needs_excel_guard(value: str) -> bool:
    """Whether Excel would evaluate ``value`` as a formula when it opens a CSV: it starts with ``=``, ``+``,
    ``@``, a tab or CR – or with ``-``, except a plain negative measurement (``-55°C``)."""
    if not value or value[0] not in _FORMULA_START:
        return False
    return not (value[0] == "-" and _SAFE_NEGATIVE.match(value))


def _read_csv(path: Path, check: Callable[[], None] | None = None) -> tuple[list[list[Any]], str, str]:
    text, encoding = _decode(path.read_bytes())
    delimiter = "\t" if path.suffix.lower() == ".tsv" else _sniff_delimiter(text)
    rows: list[list[Any]] = []
    if csv.field_size_limit() < _CSV_FIELD_LIMIT:  # long descriptions/notes (PartForge stores TEXT without limit)
        csv.field_size_limit(_CSV_FIELD_LIMIT)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        for i, r in enumerate(reader):
            if check is not None and i % _CHECK_EVERY == 0:
                check()
            rows.append([_unguard(v) for v in r])
    except csv.Error as exc:
        raise SpreadsheetError(
            f"Die CSV-Datei „{path.name}“ ist ab Zeile {reader.line_num} nicht lesbar ({exc}). Bitte die Datei prüfen "
            "(z. B. ein nicht geschlossenes Anführungszeichen) oder als XLSX speichern."
        ) from exc
    return rows, encoding, delimiter


def _percent_text(number: float) -> str:
    """A percent-formatted cell value (``0.01`` shown as ``1 %``) as text ``"1%"``."""
    return f"{to_text(round(number * 100, 10))}%"


def _is_percent_format(number_format: str | None) -> bool:
    if not number_format:
        return False
    return "%" in _FORMAT_LITERALS.sub("", number_format)


def _read_xlsx(path: Path, sheet: str | None,
               check: Callable[[], None] | None = None) -> tuple[list[str], str, list[list[Any]]]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for broken files
        raise SpreadsheetError(f"Die Excel-Datei konnte nicht gelesen werden: {exc}") from exc
    try:
        names = [str(name) for name in wb.sheetnames]
        if not names:
            raise SpreadsheetError("Die Excel-Datei enthält keine Tabellenblätter.")
        chosen = sheet if sheet is not None and sheet in names else names[0]
        ws = wb[chosen]
        rows: list[list[Any]] = []
        for i, cells in enumerate(ws.iter_rows()):
            if i >= _MAX_ROWS:
                break
            if check is not None and i % _CHECK_EVERY == 0:
                check()
            values: list[Any] = []
            for cell in cells[:_MAX_COLUMNS]:
                value = getattr(cell, "value", None)
                if (isinstance(value, (int, float)) and not isinstance(value, bool)
                        and _is_percent_format(getattr(cell, "number_format", None))):
                    value = _percent_text(value)
                values.append(value)
            rows.append(values)
        return names, chosen, rows
    finally:
        wb.close()


def _ods_cell_value(cell: Any) -> Any:
    from odf import teletype

    value_type = cell.getAttribute("valuetype")
    if value_type in ("float", "percentage", "currency"):
        raw = cell.getAttribute("value")
        try:
            number = float(raw)
        except (TypeError, ValueError):
            return teletype.extractText(cell) or None
        if value_type == "percentage":  # stored as 0.01 for "1 %"
            return _percent_text(number)
        return int(number) if number.is_integer() and abs(number) < 1e15 else number
    if value_type == "date":
        return cell.getAttribute("datevalue") or teletype.extractText(cell) or None
    if value_type == "boolean":
        return (cell.getAttribute("booleanvalue") or "").lower() == "true"
    text = teletype.extractText(cell)
    return text if text != "" else None


def _read_ods(path: Path, sheet: str | None,
              check: Callable[[], None] | None = None) -> tuple[list[str], str, list[list[Any]]]:
    from odf.namespaces import TABLENS
    from odf.opendocument import load
    from odf.table import Table, TableRow

    try:
        doc = load(str(path))
    except Exception as exc:  # noqa: BLE001
        raise SpreadsheetError(f"Die ODS-Datei konnte nicht gelesen werden: {exc}") from exc
    tables = doc.spreadsheet.getElementsByType(Table)
    names = [str(t.getAttribute("name")) for t in tables]
    if not names:
        raise SpreadsheetError("Die ODS-Datei enthält keine Tabellenblätter.")
    chosen = sheet if sheet in names else names[0]
    table = tables[names.index(chosen)]
    cell_names = {(TABLENS, "table-cell"), (TABLENS, "covered-table-cell")}
    rows: list[list[Any]] = []
    for index, tr in enumerate(table.getElementsByType(TableRow)):
        if check is not None and index % _CHECK_EVERY == 0:
            check()
        cells: list[Any] = []
        for cell in tr.childNodes:
            if getattr(cell, "qname", None) not in cell_names:
                continue
            repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
            value = _ods_cell_value(cell)
            if value is None and len(cells) + repeat > _MAX_COLUMNS:
                repeat = max(0, _MAX_COLUMNS - len(cells))
            cells.extend([value] * min(repeat, _MAX_COLUMNS))
            if len(cells) >= _MAX_COLUMNS:
                break
        while cells and cells[-1] is None:
            cells.pop()
        repeat_rows = int(tr.getAttribute("numberrowsrepeated") or 1)
        if not cells:
            repeat_rows = 1  # trailing empty rows are repeated a million times in ODS files
        rows.extend([list(cells) for _ in range(min(repeat_rows, _MAX_ROWS))])
        if len(rows) >= _MAX_ROWS:
            break
    return names, chosen, rows


def _trim(rows: list[list[Any]]) -> list[list[Any]]:
    def empty(v: Any) -> bool:
        return v is None or (isinstance(v, str) and not v.strip())

    while rows and all(empty(v) for v in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for i in range(len(row) - 1, -1, -1):
            if not empty(row[i]):
                width = max(width, i + 1)
                break
    return [list(row[:width]) + [None] * (width - len(row[:width])) for row in rows]


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _looks_numeric(value: Any) -> bool:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return True
    text = to_text(value)
    if text is None:
        return False
    return bool(re.fullmatch(r"[-+±]?\d+([.,]\d+)?\s*[a-zA-Zµμ%Ω°/]{0,6}", text))


def _shape(value: Any) -> str:
    """Character-class pattern of a cell (``CMP-428-000063`` -> ``A-9-9``) to compare rows."""
    text = to_text(value) or ""
    text = re.sub(r"[^\W\d_]+", "A", text)
    return re.sub(r"\d+", "9", text)


def _looks_like_data_row(first: list[Any], following: list[Any]) -> bool:
    """True if the first row has the same cell shapes as the row below and mostly digits –
    an export without header row (e.g. Altium 365: ``CMP-428-000063;CAP CER 68pF …``)."""
    pairs = [(a, b) for a, b in zip(first, following) if to_text(a) and to_text(b)]
    if not pairs:
        return False
    same = sum(1 for a, b in pairs if _shape(a) == _shape(b))
    with_digits = sum(1 for a, _ in pairs if any(ch.isdigit() for ch in to_text(a) or ""))
    return same / len(pairs) >= 0.6 and with_digits / len(pairs) >= 0.5


def _detect_header(rows: list[list[Any]], cat: CategoryDef | None) -> int | None:
    """0-based index of the header row or None if the file has no header row."""
    first = next((i for i, row in enumerate(rows) if any(to_text(v) for v in row)), None)
    if first is None:
        return None
    cells = [text for text in (to_text(v) for v in rows[first]) if text]
    if not cells:
        return None
    numeric = sum(1 for v in rows[first] if _looks_numeric(v))
    if numeric / len(cells) >= 0.5:
        return None
    following = next((row for row in rows[first + 1:] if any(to_text(v) for v in row)), None)
    if following is not None and _looks_like_data_row(rows[first], following):
        return None
    if cat is not None:
        scores = [score for _, score in (_best_field(c, cat) for c in cells)]
        if any(s >= MATCH_THRESHOLD for s in scores):
            return first
    if len(set(cells)) == len(cells) and all(len(c) <= 40 for c in cells):
        return first
    return None


def read_table_file(
    path: Path,
    *,
    sheet: str | None = None,
    header_row: int | Literal["auto"] | None = "auto",
    category: CategoryDef | None = None,
    check: Callable[[], None] | None = None,
) -> SheetData:
    """Read CSV/TSV/TXT, XLSX/XLSM or ODS. ``header_row``: 1-based row number,
    ``None``/0 = no header (columns are named "Spalte A", …), ``"auto"`` = detect.
    ``check`` is called every few thousand rows while reading (the import job's ``check_cancelled``)."""
    path = Path(path)
    ext = path.suffix.lower()
    encoding: str | None = None
    delimiter: str | None = None
    if ext in (".csv", ".txt", ".tsv"):
        raw_rows, encoding, delimiter = _read_csv(path, check)
        sheets: list[str] = []
        chosen: str | None = None
    elif ext in (".xlsx", ".xlsm"):
        sheets, chosen, raw_rows = _read_xlsx(path, sheet, check)
    elif ext == ".ods":
        sheets, chosen, raw_rows = _read_ods(path, sheet, check)
    else:
        raise SpreadsheetError(f"Dateityp „{ext or path.name}“ wird nicht unterstützt (CSV, XLSX oder ODS).")
    rows = _trim(raw_rows[:_MAX_ROWS])
    if header_row == "auto":
        header_index = _detect_header(rows, category)
    elif header_row:
        header_index = int(header_row) - 1
        if not 0 <= header_index < max(len(rows), 1):
            raise SpreadsheetError(f"Kopfzeile {header_row} gibt es in der Datei nicht.")
    else:
        header_index = None
    width = max((len(r) for r in rows), default=0)
    if header_index is None:
        headers = [f"Spalte {_column_letter(i)}" for i in range(width)]
        data = rows
    else:
        headers = []
        for i in range(width):
            text = to_text(rows[header_index][i]) if i < len(rows[header_index]) else None
            headers.append(text or f"Spalte {_column_letter(i)}")
        data = rows[header_index + 1:]
    first_line = 1 if header_index is None else header_index + 2
    kept = [(first_line + i, r) for i, r in enumerate(data) if any(to_text(v) for v in r)]
    data = [r for _, r in kept]
    return SheetData(
        sheets=sheets,
        sheet=chosen,
        headers=headers,
        rows=data,
        header_row=None if header_index is None else header_index + 1,
        encoding=encoding,
        delimiter=delimiter,
        raw_first_rows=rows[:5],
        lines=[line for line, _ in kept],
    )


# --------------------------------------------------------------------------- mapping


class MappingItem(BaseModel):
    """Source column (by index) -> target field (None = ignore).

    ``unit_hint``: unit for plain numbers of this column (from the header, e.g. ``pF`` in
    ``Kapazität (pF)``); ``100`` is imported as ``100pF``. Not sent = derived from the
    header again; ``null``/``""`` = no unit hint. ``unit_note`` explains it in the preview."""

    index: int
    header: str = ""
    column: str | None = None
    score: float = 0.0
    unit_hint: str | None = Field(default=None, description="Einheit für Zahlen ohne Einheit, z. B. „pF“")
    unit_note: str | None = Field(default=None, description="Hinweis zur Einheit (nur Vorschau)")


def _normalize_header(text: str) -> str:
    return utils.default_process(text.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
                                 .replace("Ä", "ae").replace("Ö", "oe").replace("Ü", "ue"))


def _choices(cat: CategoryDef) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for f in cat.fields:
        for text in (f.column, f.label, *SYNONYMS.get(f.column, [])):
            normalized = _normalize_header(text)
            if normalized:
                out.append((normalized, f.column))
    return out


_choice_cache: dict[tuple[str, int], list[tuple[str, str]]] = {}


def _best_field(header: str, cat: CategoryDef) -> tuple[str | None, float]:
    key = (cat.table, hash(tuple((f.column, f.label) for f in cat.fields)))
    choices = _choice_cache.get(key)
    if choices is None:
        choices = _choices(cat)
        _choice_cache[key] = choices
    variants = {_normalize_header(header)}
    stripped = re.sub(r"[\(\[].*?[\)\]]", " ", header)
    variants.add(_normalize_header(stripped))
    variants.discard("")
    best: tuple[str | None, float] = (None, 0.0)
    for variant in variants:
        for normalized, column in choices:
            if variant == normalized:
                return column, 100.0
            score = max(fuzz.ratio(variant, normalized), fuzz.token_sort_ratio(variant, normalized))
            if score > best[1]:
                best = (column, float(score))
    return best


def _hint_scale(hint: str, field: FieldDef) -> tuple[bool, str | None]:
    """(compatible, parsed unit) of a unit hint for a quantity field (``kOhm`` for Ω)."""
    q = parse_field_quantity(f"1{hint}", field.unit)
    if q is None:
        return False, None
    return units_compatible(q.unit, field.unit), q.unit


def header_unit(header: str, field: FieldDef) -> tuple[str | None, str | None]:
    """``(unit_hint, note)`` from a unit in brackets in a column header.

    ``"Kapazität (pF)"`` for a capacitance field -> ``("pF", "Zahlen ohne Einheit werden als pF gelesen.")``;
    an incompatible unit (``"Kapazität (V)"``) gives no hint but a German warning; words in
    brackets that are no unit (``"(typ)"``) are ignored."""
    if field.kind not in ("quantity", "range") or not field.unit:
        return None, None
    for match in _HEADER_UNIT.finditer(header or ""):
        hint = match.group(1).strip()
        if not hint:
            continue
        compatible, parsed_unit = _hint_scale(hint, field)
        if parsed_unit is None and not compatible:
            continue  # "(typ)", "(max)", "(Stück)": not a unit
        if not compatible:
            return None, (
                f"Die Einheit „{hint}“ im Spaltenkopf passt nicht zum Feld „{field.label}“ ({field.unit}) – "
                f"Zahlen ohne Einheit werden als {field.unit} gelesen. Bitte die Zuordnung prüfen."
            )
        return hint, f"Zahlen ohne Einheit werden als {hint} gelesen."
    return None, None


def apply_unit_hint(value: Any, hint: str | None, kind: str = "quantity") -> Any:
    """Append the unit hint to a plain number (``100`` -> ``"100pF"``, ``"4,7"`` -> ``"4,7kOhm"``).

    Values with their own unit or prefix (``"100n"``, ``"4k7"``) and non-numbers stay
    unchanged; for ranges each plain end gets the hint (``"-40 … 85"`` -> ``"-40°C … 85°C"``)."""
    if not hint or value is None or isinstance(value, bool):
        return value
    text = to_text(value)
    if text is None:
        return value
    if _PLAIN_NUMBER.match(text):
        return f"{text}{hint}"
    if kind == "range":
        pieces = _RANGE_PARTS.split(text)
        if len(pieces) == 3 and all(_PLAIN_NUMBER.match(p.strip()) for p in (pieces[0], pieces[2])):
            return f"{pieces[0].strip()}{hint}{pieces[1]}{pieces[2].strip()}{hint}"
    return value


def suggest_mapping(headers: Sequence[str], cat: CategoryDef) -> list[MappingItem]:
    """Best field per source column (one-to-one, highest scores first) incl. unit hints."""
    candidates = [(i, h, *_best_field(h, cat)) for i, h in enumerate(headers)]
    taken: set[str] = set()
    result = {i: MappingItem(index=i, header=h) for i, h in enumerate(headers)}
    for i, header, column, score in sorted(candidates, key=lambda c: -c[3]):
        if column is None or score < MATCH_THRESHOLD or column in taken:
            continue
        taken.add(column)
        f = cat.field(column)
        hint, note = header_unit(header, f) if f is not None else (None, None)
        result[i] = MappingItem(index=i, header=header, column=column, score=round(score, 1), unit_hint=hint,
                                unit_note=note)
    return [result[i] for i in range(len(headers))]


def _validate_mapping(mapping: Sequence[MappingItem | Mapping[str, Any]], headers: Sequence[str], cat: CategoryDef) -> list[MappingItem]:
    items = [m if isinstance(m, MappingItem) else MappingItem.model_validate(dict(m)) for m in mapping]
    used: dict[str, int] = {}
    clean: list[MappingItem] = []
    for item in items:
        if not item.column:
            continue
        if not 0 <= item.index < len(headers):
            raise SpreadsheetError(f"Spalte {item.index + 1} gibt es in der Datei nicht.")
        f = cat.field(item.column)
        if f is None:
            raise SpreadsheetError(f"Das Feld „{item.column}“ gibt es in der Kategorie „{cat.label}“ nicht.")
        if f.column in used:
            raise SpreadsheetError(f"Das Feld „{f.column}“ ist mehreren Spalten zugeordnet ({headers[used[f.column]]}, {headers[item.index]}).")
        used[f.column] = item.index
        header = headers[item.index]
        if "unit_hint" not in item.model_fields_set:  # client did not decide: derive from the header
            hint, note = header_unit(header, f)
        else:
            hint, note = (item.unit_hint or "").strip() or None, item.unit_note
            if hint is not None:
                if f.kind not in ("quantity", "range") or not f.unit:
                    raise SpreadsheetError(f"Das Feld „{f.label}“ hat keine Einheit – „{hint}“ (Spalte „{header}“) passt nicht.")
                if not _hint_scale(hint, f)[0]:
                    raise SpreadsheetError(
                        f"Die Einheit „{hint}“ (Spalte „{header}“) passt nicht zum Feld „{f.label}“ ({f.unit})."
                    )
        clean.append(item.model_copy(update={"column": f.column, "header": header, "unit_hint": hint, "unit_note": note}))
    if not clean:
        raise SpreadsheetError("Keine Spalte ist einem Feld zugeordnet.")
    return clean


def build_rows(
    data: SheetData,
    mapping: Sequence[MappingItem],
    defaults: Mapping[str, Any] | None = None,
    category: CategoryDef | None = None,
) -> list[dict[str, Any]]:
    """Rows as ``{field: value}`` (mapping and unit hints applied; defaults fill empty cells)."""
    defaults = {k: v for k, v in (defaults or {}).items() if to_text(v) is not None}
    kinds: dict[int, str] = {}
    for item in mapping:
        f = category.field(str(item.column)) if category is not None and item.column else None
        kinds[item.index] = f.kind if f is not None else "quantity"
    out: list[dict[str, Any]] = []
    for raw in data.rows:
        row: dict[str, Any] = dict(defaults)
        for item in mapping:
            value = raw[item.index] if item.index < len(raw) else None
            if item.unit_hint:
                value = apply_unit_hint(value, item.unit_hint, kinds[item.index])
            if to_text(value) is not None or item.column not in row:
                row[str(item.column)] = value
        out.append(row)
    return out


# --------------------------------------------------------------------------- uploads


def _upload_dir() -> Path:
    """Uploaded import files (kept 24 h): in the local, non-roaming data folder."""
    folder = local_data_dir() / "import"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _cleanup_uploads() -> None:
    cutoff = time.time() - UPLOAD_MAX_AGE_S
    for file in _upload_dir().iterdir():
        try:
            if file.is_file() and file.stat().st_mtime < cutoff:
                file.unlink()
        except OSError:  # pragma: no cover - best effort
            log.debug("Alte Importdatei nicht löschbar: %s", file)


def store_upload(data: bytes, filename: str) -> str:
    """Keep an uploaded file for preview/run; returns its token."""
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise SpreadsheetError(f"Dateityp „{ext or filename}“ wird nicht unterstützt (CSV, XLSX oder ODS).")
    if not data:
        raise SpreadsheetError("Die Datei ist leer.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise SpreadsheetError(f"Die Datei ist zu groß (höchstens {MAX_UPLOAD_BYTES // (1024 * 1024)} MB).")
    _cleanup_uploads()
    token = secrets.token_hex(12)
    (_upload_dir() / f"{token}{ext}").write_bytes(data)
    (_upload_dir() / f"{token}.name").write_text(Path(filename).name, encoding="utf-8")
    return token


def upload_path(token: str) -> tuple[Path, str]:
    """Stored file and original name of an upload token."""
    if not _TOKEN_RE.match(token or ""):
        raise SpreadsheetError("Ungültiger Upload-Schlüssel – bitte die Datei erneut hochladen.")
    for ext in SUPPORTED_EXTENSIONS:
        candidate = _upload_dir() / f"{token}{ext}"
        if candidate.is_file():
            name_file = _upload_dir() / f"{token}.name"
            name = name_file.read_text(encoding="utf-8") if name_file.is_file() else candidate.name
            return candidate, name
    raise SpreadsheetError("Die hochgeladene Datei ist nicht mehr vorhanden – bitte erneut hochladen.")


# --------------------------------------------------------------------------- per-row status (dry run)

ImportRowStatusName = Literal["create", "update", "unchanged", "skip", "reject"]
ImportMode = Literal["skip", "update", "error"]
CHECK_ROWS_DEFAULT = 200
CHECK_ROWS_MAX = 5000


class ImportConflict(BaseModel):
    table: str
    part_number: str


class ImportRowStatus(BaseModel):
    """What the import would do with one data row (dry run with the current mapping)."""

    row: int = Field(description="Index der Datenzeile (0-basiert, wie in ImportPreviewOut.rows)")
    line: int | None = Field(default=None, description="Zeile in der Datei/Tabelle (1-basiert)")
    part_number: str | None = Field(default=None, description="Teilenummer (None: wird beim Import vergeben)")
    status: ImportRowStatusName = Field(
        description="create = neu, update = aktualisieren, unchanged = vorhanden ohne Änderung, "
                    "skip = vorhanden (Modus skip), reject = abgelehnt"
    )
    reasons: list[str] = Field(default_factory=list, description="Gründe (abgelehnt/übersprungen), deutsch")
    errors: dict[str, str] = Field(default_factory=dict, description="Fehler je Feld (Spaltenname → Meldung)")
    warnings: list[str] = Field(default_factory=list)
    missing: list[str] = Field(default_factory=list, description="Leere Pflichtfelder (Import als Entwurf)")
    changed: list[str] = Field(default_factory=list, description="Felder, die ein update ändert")
    conflict: ImportConflict | None = Field(default=None, description="Bauteil, dem die Teilenummer gehört")


class ImportCheck(BaseModel):
    """Summary of the per-row dry run."""

    checked: int = Field(description="Geprüfte Zeilen (die ersten n)")
    total: int = Field(description="Datenzeilen der Datei")
    truncated: bool = Field(description="Nicht alle Zeilen geprüft")
    counts: dict[str, int] = Field(description="Zeilen je Status (create, update, unchanged, skip, reject)")
    mode: ImportMode
    strict: bool = False
    allow_incomplete: bool = False


def truncated_columns(row: Mapping[str, Any]) -> list[str]:
    """Fields of an input row whose text has exactly :data:`XLSX_CELL_LIMIT` characters – most likely cut by
    Excel (an XLSX cell holds no more)."""
    return [str(k) for k, v in row.items() if isinstance(v, str) and len(v) == XLSX_CELL_LIMIT]


def _truncation_note(columns: Sequence[str]) -> str:
    each = "je " if len(columns) > 1 else ""
    return (f"{', '.join(columns)}: {each}genau {_CELL_LIMIT_TEXT} Zeichen – vermutlich von Excel gekürzt (eine "
            "XLSX-Zelle fasst nicht mehr). Längere Texte als CSV übertragen.")


def _keep_longer_stored(table: str, rows: list[dict[str, Any]], mode: str) -> dict[int, list[str]]:
    """Mode ``update``: a cell of exactly :data:`XLSX_CELL_LIMIT` characters that is the start of the longer
    stored text (an XLSX round trip) never overwrites it – the field is dropped from the row. Returns
    ``{row index: [columns kept]}``."""
    kept: dict[int, list[str]] = {}
    if mode != "update":
        return kept
    for index, row in enumerate(rows):
        columns = truncated_columns(row)
        if not columns:
            continue
        pn = to_text(next((v for k, v in row.items() if str(k).casefold() == PART_NUMBER.casefold()), None))
        found = parts.find_part(pn) if pn else None
        if found is None or found[0].casefold() != table.casefold():
            continue
        stored = found[1]
        for column in columns:
            current = next((v for k, v in stored.items() if k.casefold() == column.casefold()), None)
            text = str(row[column])
            if isinstance(current, str) and len(current) > len(text) and current.startswith(text.rstrip()):
                del row[column]
                kept.setdefault(index, []).append(column)
    return kept


def _full_mapping(headers: Sequence[str], mapped: Sequence[MappingItem]) -> list[MappingItem]:
    """One item per source column: the mapped ones, the others without target field."""
    by_index = {m.index: m for m in mapped}
    return [by_index.get(i, MappingItem(index=i, header=h)) for i, h in enumerate(headers)]


def check_rows(
    data: SheetData,
    cat: CategoryDef,
    mapping: Sequence[MappingItem],
    *,
    mode: ImportMode = "skip",
    defaults: Mapping[str, Any] | None = None,
    allow_incomplete: bool = False,
    strict: bool = False,
    limit: int = CHECK_ROWS_DEFAULT,
    thousands: bool | None = None,
) -> tuple[list[ImportRowStatus], ImportCheck]:
    """Dry run of the import for the first ``limit`` rows (validated ``mapping``): per-row status and
    counts. Nothing is written."""
    limit = max(0, min(int(limit), CHECK_ROWS_MAX))
    window = SheetData(data.sheets, data.sheet, data.headers, data.rows[:limit], data.header_row)
    rows = build_rows(window, mapping, defaults, category=cat) if mapping else []
    cut = {i: truncated_columns(row) for i, row in enumerate(rows)}
    kept = _keep_longer_stored(cat.table, rows, mode)
    plans = parts.plan_upsert(cat.table, rows, mode=mode, allow_incomplete=allow_incomplete, strict=strict,
                              thousands=thousands)
    statuses: list[ImportRowStatus] = []
    counts = dict.fromkeys(("create", "update", "unchanged", "skip", "reject"), 0)
    for plan in plans:
        if cut.get(plan.index):
            plan.warnings = [_truncation_note(cut[plan.index]), *plan.warnings]
        if kept.get(plan.index):
            plan.warnings = [f"{', '.join(kept[plan.index])}: der gespeicherte, längere Text bleibt erhalten.",
                             *plan.warnings]
        counts[plan.status] += 1
        statuses.append(ImportRowStatus(
            row=plan.index,
            line=data.lines[plan.index] if plan.index < len(data.lines) else None,
            part_number=plan.part_number,
            status=plan.status,
            reasons=plan.reasons,
            errors=plan.errors,
            warnings=plan.warnings,
            missing=plan.missing,
            changed=plan.changed,
            conflict=ImportConflict(table=plan.conflict[0], part_number=plan.conflict[1]) if plan.conflict else None,
        ))
    check = ImportCheck(checked=len(rows), total=len(data.rows), truncated=len(rows) < len(data.rows), counts=counts,
                        mode=mode, strict=strict, allow_incomplete=allow_incomplete)
    return statuses, check


# --------------------------------------------------------------------------- preview & run


def preview(
    path: Path,
    *,
    table: str | None = None,
    sheet: str | None = None,
    header_row: int | Literal["auto"] | None = "auto",
    sample: int = 20,
    mapping: Sequence[MappingItem | Mapping[str, Any]] | None = None,
    mode: ImportMode = "skip",
    defaults: Mapping[str, Any] | None = None,
    allow_incomplete: bool = False,
    strict: bool = False,
    check_limit: int = CHECK_ROWS_DEFAULT,
    thousands: bool | None = None,
) -> dict[str, Any]:
    """Sheets, headers, mapping, sample rows (+ status new/exists per sample row) and – with a
    category – the per-row dry run (``row_status``, ``check``) for ``mapping`` (``None`` = the
    suggested one), ``mode``, ``defaults``, ``allow_incomplete``, ``strict`` and ``thousands``."""
    cat = categories.get_category(table) if table else None
    data = read_table_file(path, sheet=sheet, header_row=header_row, category=cat)
    result: dict[str, Any] = {
        "sheets": data.sheets,
        "sheet": data.sheet,
        "headers": data.headers,
        "header_row": data.header_row,
        "encoding": data.encoding,
        "delimiter": data.delimiter,
        "row_count": len(data.rows),
        "rows": [[to_text(v) for v in row] for row in data.rows[:sample]],
        "mapping": [],
        "samples": [],
        "row_status": [],
        "check": None,
    }
    if cat is None:
        return result
    if mapping is None:
        chosen = suggest_mapping(data.headers, cat)
        mapped = [m for m in chosen if m.column]
    else:
        items = [m if isinstance(m, MappingItem) else MappingItem.model_validate(dict(m)) for m in mapping]
        mapped = _validate_mapping(items, data.headers, cat) if any(m.column for m in items) else []
        chosen = _full_mapping(data.headers, mapped)
    result["mapping"] = [m.model_dump() for m in chosen]
    if mapped:
        statuses, check = check_rows(data, cat, mapped, mode=mode, defaults=defaults, allow_incomplete=allow_incomplete,
                                     strict=strict, limit=check_limit, thousands=thousands)
        result["row_status"] = [st.model_dump() for st in statuses]
        result["check"] = check.model_dump()
    if mapped:
        sample_rows = build_rows(SheetData(data.sheets, data.sheet, data.headers, data.rows[:sample], data.header_row),
                                 mapped, defaults, category=cat)
        pns = [to_text(r.get(PART_NUMBER)) or to_text(r.get("Manufacturer Part Number")) for r in sample_rows]
        existing = parts.existing_part_numbers([p for p in pns if p])
        samples = []
        for row, pn in zip(sample_rows, pns):
            status = "neu" if pn and pn not in existing else ("vorhanden" if pn else "ohne Teilenummer")
            samples.append({"part_number": pn, "status": status, "values": {k: to_text(v) for k, v in row.items()}})
        result["samples"] = samples
    return result


class ImportOptions(BaseModel):
    """Parameters of an import run (API body of ``POST /import/run``)."""

    token: str
    table: str
    mapping: list[MappingItem]
    mode: Literal["skip", "update", "error"] = "skip"
    sheet: str | None = None
    header_row: int | Literal["auto"] | None = "auto"
    defaults: dict[str, Any] = Field(default_factory=dict)  # constant values, e.g. {"Manufacturer": "Yageo"}
    allow_incomplete: bool = False  # import rows with empty required fields as drafts
    strict: bool = Field(default=False, description="Nicht lesbare Werte (keine Zahl, falsche Einheit, kein Datum …) "
                                                     "lehnen die Zeile ab statt nur zu warnen")
    thousands: bool | None = Field(
        default=None,
        description="Zahlenformat der Quelle für mehrdeutige Kommas wie „1,000“: None = unbekannt (Warnung, im "
                    "strengen Modus abgelehnt), true = Komma als Tausendertrennzeichen (englische Quelle: 1000), "
                    "false = Dezimalkomma (deutsche Quelle: 1,0)")
    author: str | None = None


ImportOutcome = Literal["ok", "partial", "failed"]

#: How an import treats rows whose part number exists already (job log, German).
MODE_LABELS: dict[str, str] = {
    "skip": "vorhandene überspringen",
    "update": "vorhandene aktualisieren",
    "error": "vorhandene ablehnen",
}
_LOGGED_ERRORS = 50  # rejected rows listed in the job log (the result keeps 200)


def import_outcome(error_count: int, written: int) -> ImportOutcome:
    """``ok`` without rejected rows, ``partial`` if parts were created/updated despite rejected rows,
    ``failed`` if rows were rejected and nothing was created or updated."""
    if not error_count:
        return "ok"
    return "partial" if written else "failed"


def import_summary(*, created: int, updated: int, unchanged: int = 0, skipped: int = 0, rejected: int = 0) -> str:
    """Counts of an import run, zero counts left out: „120 neu, 4 aktualisiert, 2 abgelehnt“
    (thousands with a dot: „1.250 neu“); without any row: „keine Zeile“."""
    counts = ((created, "neu"), (updated, "aktualisiert"), (unchanged, "unverändert"), (skipped, "übersprungen"),
              (rejected, "abgelehnt"))
    return ", ".join(plural(n, word, word) for n, word in counts if n) or "keine Zeile"


def import_message(outcome: ImportOutcome, summary: str) -> str:
    """Final message of the import job: „Import fertig: 120 neu, 4 aktualisiert, 2 abgelehnt“ (``ok`` and
    ``partial``) or „Import fehlgeschlagen – nichts übernommen: 2 abgelehnt“ (``failed``)."""
    if outcome == "failed":
        return f"Import fehlgeschlagen – nichts übernommen: {summary}"
    return f"Import fertig: {summary}"


def _cancel_note(ctx: JobContext, exc: JobCancelled) -> str:
    """Job log line of a cancelled import (the parts are written in one transaction: nothing was imported)."""
    if isinstance(exc, LibraryChanged):
        return str(exc)
    if cancelled_by_switch(ctx):
        return "Die Bibliothek wird geschlossen oder gewechselt – Import abgebrochen, es wurde nichts übernommen."
    return "Import abgebrochen – es wurde nichts übernommen."


def run_import(
    ctx: JobContext | None,
    options: ImportOptions | Mapping[str, Any],
    *,
    library: str | os.PathLike[str] | None = None,
) -> dict[str, Any]:
    """Import job: read file, map, backup, :func:`upsert_parts`. Returns a JSON-able summary (with the
    final ``message``).

    ``library``: root folder of the library the import was started for (API: the open library at
    submit time). The job then holds it for its whole run (:func:`partforge.jobs.library_hold`, job kind
    ``import`` in :data:`partforge.jobs.LIBRARY_JOB_KINDS`): a library switch cancels the import and waits
    until it stopped – the parts are written in one transaction, so a cancelled import imported nothing;
    if that library is no longer open when the job starts, it stops (``LibraryChanged``) without reading
    or writing anything. ``None`` (direct calls): the open library, no hold.
    """
    opts = options if isinstance(options, ImportOptions) else ImportOptions.model_validate(dict(options))
    if library is None:
        return _run_import(ctx, opts)
    try:
        with library_hold(ctx, library, announce=False):  # waits while a library switch runs
            return _run_import(ctx, opts)
    except JobCancelled as exc:
        if ctx is not None:
            ctx.log(_cancel_note(ctx, exc), "warn")
        raise


def _run_import(ctx: JobContext | None, opts: ImportOptions) -> dict[str, Any]:
    path, original = upload_path(opts.token)
    cat = categories.get_category(opts.table)
    if ctx is not None:
        ctx.progress(0.02, f"„{original}“ wird gelesen …")
    data = read_table_file(path, sheet=opts.sheet, header_row=opts.header_row, category=cat,
                           check=ctx.check_cancelled if ctx is not None else None)
    mapping = _validate_mapping(opts.mapping, data.headers, cat)
    rows = build_rows(data, mapping, opts.defaults, category=cat)
    if ctx is not None:
        for item in mapping:
            if item.unit_hint:
                ctx.log(f"Spalte „{item.header}“: Zahlen ohne Einheit werden als {item.unit_hint} gelesen.")
    if not rows:
        raise SpreadsheetError("Die Datei enthält keine Datenzeilen.")
    if ctx is not None:
        ctx.check_cancelled()  # nothing written yet
        ctx.log(f"{plural(len(rows), 'Zeile', 'Zeilen')} aus „{original}“ → {cat.label} "
                f"(Modus: {MODE_LABELS.get(opts.mode, opts.mode)})")
    backups: list[Path] = []

    def backup_first() -> None:  # only when rows will be written: an import that changes nothing needs none
        if ctx is not None:
            ctx.check_cancelled()
            ctx.progress(None, "Sicherung wird angelegt …")
        made = backup.auto_backup("vor_import")
        if made is not None:
            backups.append(made)
            if ctx is not None:
                ctx.log(f"Sicherung angelegt: {made.name}")
    # the last cancellation point: upsert_parts writes everything in one transaction (rolled back when cancelled)
    cut = sum(1 for row in rows if truncated_columns(row))
    kept = _keep_longer_stored(cat.table, rows, opts.mode)
    if ctx is not None and cut:
        ctx.log(f"{plural(cut, 'Zeile hat', 'Zeilen haben')} Zellen mit genau {_CELL_LIMIT_TEXT} Zeichen – "
                "vermutlich von Excel gekürzt.", "warn")
    if ctx is not None and kept:
        ctx.log(f"{plural(len(kept), 'Bauteil behält', 'Bauteile behalten')} den gespeicherten, längeren Text "
                "(gekürzte Zellen nicht übernommen).", "warn")
    if ctx is not None and opts.thousands is not None:
        ctx.log("Zahlen wie „1,000“ werden mit Komma als " + ("Tausendertrennzeichen (1000)" if opts.thousands
                                                               else "Dezimalkomma (1,0)") + " gelesen.")
    result = parts.upsert_parts(cat.table, rows, mode=opts.mode, author=opts.author, ctx=ctx,
                                allow_incomplete=opts.allow_incomplete, strict=opts.strict, thousands=opts.thousands,
                                before_write=backup_first)
    backup_file = backups[0] if backups else None
    outcome = import_outcome(len(result.errors), len(result.created) + len(result.updated))
    summary = import_summary(created=len(result.created), updated=len(result.updated),
                             unchanged=len(result.unchanged), skipped=len(result.skipped), rejected=len(result.errors))
    message = import_message(outcome, summary)
    if ctx is not None:
        for pn, text in list(result.errors.items())[:_LOGGED_ERRORS]:
            ctx.log(f"{pn}: {text}", "warn")
        if len(result.errors) > _LOGGED_ERRORS:
            more = len(result.errors) - _LOGGED_ERRORS
            ctx.log(f"… und {plural(more, 'weitere abgelehnte Zeile', 'weitere abgelehnte Zeilen')}", "warn")
        ctx.log(message, "success" if outcome == "ok" else "warn")
        ctx.progress(1.0, message)
    return {
        "outcome": outcome,
        "message": message,
        "strict": opts.strict,
        "table": cat.table,
        "file": original,
        "rows": len(rows),
        "created": len(result.created),
        "updated": len(result.updated),
        "unchanged": len(result.unchanged),
        "skipped": len(result.skipped),
        "error_count": len(result.errors),
        "errors": dict(list(result.errors.items())[:200]),
        "warnings": dict(list(result.warnings.items())[:100]),
        "created_part_numbers": result.created[:500],
        "backup": backup_file.name if backup_file else None,
    }


# --------------------------------------------------------------------------- export


def _export_rows(
    cat: CategoryDef,
    *,
    pns: Iterable[str] | None,
    search: str | None,
    filters: Sequence[PartFilter | Mapping[str, Any]] | None,
    sort: str | Sequence[SortSpec | Mapping[str, Any] | str] | None,
) -> list[dict[str, Any]]:
    rows, _ = parts.query_parts(cat.table, search=search, filters=filters, sort=sort, limit=None)
    if pns is not None:
        wanted = {parts.pn_key(p) for p in pns if p}
        rows = [r for r in rows if parts.pn_key(str(r.get(PART_NUMBER) or "")) in wanted]
    return rows


def _export_columns(cat: CategoryDef, columns: Sequence[str] | None) -> list[str]:
    if not columns:
        return [f.column for f in cat.fields]
    out: list[str] = []
    for name in columns:
        f = cat.field(name)
        if f is None:
            raise SpreadsheetError(f"Das Feld „{name}“ gibt es in der Kategorie „{cat.label}“ nicht.")
        if f.column not in out:
            out.append(f.column)
    return out


def _safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", text).strip("_") or "Export"


@dataclass
class ExportFile:
    content: bytes
    file_name: str
    media_type: str
    rows: int
    #: XLSX: cells longer than Excel's limit (cut by the format) – reported to the user
    truncated: int = 0
    warnings: list[str] = field(default_factory=list)


def render_export(
    table: str,
    *,
    fmt: Literal["csv", "xlsx"] = "xlsx",
    pns: Iterable[str] | None = None,
    columns: Sequence[str] | None = None,
    search: str | None = None,
    filters: Sequence[PartFilter | Mapping[str, Any]] | None = None,
    sort: str | Sequence[SortSpec | Mapping[str, Any] | str] | None = None,
    delimiter: str = ";",
    excel_safe: bool = True,
    progress: Callable[[int, int], None] | None = None,
) -> ExportFile:
    """Export a category (all parts, a selection via ``pns`` or a search/filter result).

    ``progress(done, total)`` is called while the rows are written (every :data:`_PROGRESS_EVERY` rows; the
    export job reports it and may raise ``JobCancelled`` there – a large XLSX takes seconds).

    ``excel_safe`` (CSV only, default): prefix values Excel would evaluate as a formula
    (:func:`needs_excel_guard`: ``= + @``, tab, CR, ``-`` except plain negative measurements) with
    an apostrophe, so Excel shows them as text (PartForge's own import removes the apostrophe
    again); ``False`` writes every value unchanged. XLSX cells are always literal text."""
    if fmt not in ("csv", "xlsx"):
        raise SpreadsheetError(f"Unbekanntes Exportformat „{fmt}“ (csv oder xlsx).")
    if fmt == "csv" and delimiter not in _DELIMITERS:
        raise SpreadsheetError(f"Ungültiges Trennzeichen „{delimiter}“ (erlaubt: ; , Tab |).")
    cat = categories.get_category(table)
    cols = _export_columns(cat, columns)
    rows = _export_rows(cat, pns=pns, search=search, filters=filters, sort=sort)
    stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    base = f"{_safe_name(cat.table)}_{stamp}"
    if fmt == "csv":
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow([_csv_cell(c, excel_safe=excel_safe) for c in cols])
        for index, row in enumerate(rows):
            if progress is not None and index % _PROGRESS_EVERY == 0:
                progress(index, len(rows))
            writer.writerow([_csv_cell(row.get(c), excel_safe=excel_safe) for c in cols])
        return ExportFile(buffer.getvalue().encode("utf-8-sig"), f"{base}.csv", "text/csv; charset=utf-8", len(rows))
    long_cells = [(str(row.get(PART_NUMBER) or ""), col) for row in rows for col in cols
                  if isinstance(row.get(col), str) and len(row[col]) > XLSX_CELL_LIMIT]
    warnings: list[str] = []
    if long_cells:
        sample = ", ".join(f"{pn} ({col})" for pn, col in long_cells[:3]) + (" …" if len(long_cells) > 3 else "")
        warnings.append(f"{plural(len(long_cells), 'Wert', 'Werte')} länger als {_CELL_LIMIT_TEXT} Zeichen in der "
                        f"XLSX-Datei gekürzt (Excel-Grenze je Zelle): {sample}. Vollständig nur als CSV.")
        log.warning("XLSX-Export %s: %d Zellen gekürzt", cat.table, len(long_cells))
    return ExportFile(_xlsx(cat, cols, rows, progress), f"{base}.xlsx",
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", len(rows),
                      truncated=len(long_cells), warnings=warnings)


def _csv_cell(value: Any, *, excel_safe: bool) -> Any:
    if value is None:
        return ""
    if excel_safe and isinstance(value, str) and (needs_excel_guard(value) or _unguard(value) != value):
        # a formula-like value gets the guard apostrophe; a value that itself starts with "'" + "=+-@" gets
        # one more, so the importer (which strips exactly one) gets the original back
        return f"{_EXCEL_GUARD}{value}"
    return value


def export_parts(table: str, **kwargs: Any) -> tuple[bytes, str, str]:
    """Export a category; returns ``(content, file name, media type)`` (see :func:`render_export`)."""
    result = render_export(table, **kwargs)
    return result.content, result.file_name, result.media_type


def _xlsx(cat: CategoryDef, cols: list[str], rows: list[dict[str, Any]],
          progress: Callable[[int, int], None] | None = None) -> bytes:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook(write_only=True)
    title = re.sub(r"[\[\]:*?/\\]", "_", cat.table)[:31] or "Export"
    ws = wb.create_sheet(title)
    for i, col in enumerate(cols):
        sample = [len(str(r.get(col) or "")) for r in rows[:500]]
        width = max([len(col), *sample]) + 2
        ws.column_dimensions[get_column_letter(i + 1)].width = max(8, min(60, width))
    ws.freeze_panes = "A2"
    last = get_column_letter(max(len(cols), 1))
    ws.auto_filter.ref = f"A1:{last}{len(rows) + 1}"
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EBF2")
    header: list[Any] = []
    for col in cols:
        cell = WriteOnlyCell(ws, value=col)
        cell.data_type = "s"  # a column name is text, never a formula
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        header.append(cell)
    ws.append(header)

    def data_cell(value: Any) -> Any:
        if not isinstance(value, str):
            return value
        if ILLEGAL_CHARACTERS_RE.search(value):  # control characters are not allowed in XLSX
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if value[:1] in _FORMULA_START:
            # A leading = + - @ (tab, CR) must stay literal text: openpyxl writes a string starting with
            # "=" as a formula; forcing the string type keeps every such value a plain string (QA sec
            # finding 3). Altium indirection like "=Value" is a string, never evaluated.
            literal = WriteOnlyCell(ws, value=value)
            literal.data_type = "s"
            return literal
        return value

    for index, row in enumerate(rows):
        if progress is not None and index % _PROGRESS_EVERY == 0:
            progress(index, len(rows))
        ws.append([data_cell(row.get(c)) for c in cols])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_to_file(
    table: str,
    *,
    fmt: Literal["csv", "xlsx"] = "xlsx",
    pns: Iterable[str] | None = None,
    columns: Sequence[str] | None = None,
    folder: Path | None = None,
    library: str | os.PathLike[str] | None = None,
    warnings: list[str] | None = None,
    **kwargs: Any,
) -> tuple[Path, int]:
    """Export into the workspace ``Exports`` folder (or ``folder``); returns ``(path, row count)``.
    ``warnings`` (a list) receives the German notes of the export (cut XLSX cells).

    ``library`` (export job): root folder of the library the export was started for. The export only
    reads, so it holds nothing – a library switch does not wait for it. Instead it reads the parts only
    while that library is open and writes the file into *its* ``Exports`` folder only if it is still the
    open library after reading; otherwise :class:`~partforge.jobs.LibraryChanged` (a ``JobCancelled``,
    also when the library is closed while reading) – it never exports another library's parts or writes
    into another library's folder.
    """
    from ..db import core

    if library is None:
        result = render_export(table, fmt=fmt, pns=pns, columns=columns, **kwargs)
        target_dir = folder or core.current_workspace().exports_dir
    else:
        ws = open_library(library)
        if ws is None:
            raise LibraryChanged(library)
        try:
            result = render_export(table, fmt=fmt, pns=pns, columns=columns, **kwargs)
        except Exception:
            if open_library(library) is not ws:  # closed or switched while reading
                raise LibraryChanged(library) from None
            raise
        if open_library(library) is not ws:  # switched while reading: the parts may come from another library
            raise LibraryChanged(library)
        target_dir = folder or ws.exports_dir
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / result.file_name
    counter = 2
    while target.exists():
        target = target_dir / f"{Path(result.file_name).stem}-{counter}{Path(result.file_name).suffix}"
        counter += 1
    target.write_bytes(result.content)
    log.info("Export geschrieben: %s (%d Bauteile)", target, result.rows)
    if warnings is not None:
        warnings.extend(result.warnings)
    return target, result.rows
