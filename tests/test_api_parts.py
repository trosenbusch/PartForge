"""API: /parts, /search, /stats, /health-check, /distinct, /import, /export, /backups, /views."""
from __future__ import annotations

import io
import json
import time
from typing import Any
from urllib.parse import quote

import openpyxl
import pytest
from fastapi.testclient import TestClient

from partforge.db import parts as db_parts


def res(pn: str, value: str, **kw: Any) -> dict[str, Any]:
    return {"Part Number": pn, "Manufacturer": "YAGEO", "Manufacturer Part Number": pn, "Resistance": value,
            "Tolerance": "1", "Power": "0.1", "Size": "0603", "Library Ref": "RES", "Footprint Ref": "RESC1608X55N", **kw}


def url(table: str, pn: str, suffix: str = "") -> str:
    return f"/api/parts/{quote(table)}/{quote(pn, safe='')}{suffix}"


def create(client: TestClient, table: str, values: dict[str, Any], **kw: Any) -> dict[str, Any]:
    r = client.post(f"/api/parts/{quote(table)}", json={"values": values, **kw})
    assert r.status_code == 201, r.text
    return r.json()


def wait_for_job(client: TestClient, job_id: str, timeout: float = 20.0) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("done", "failed", "cancelled"):
            return job
        time.sleep(0.05)
    raise AssertionError("Job wurde nicht fertig")


@pytest.fixture
def filled(client: TestClient) -> TestClient:
    for pn, value in (("R-10K", "10k"), ("R-4K7", "4k7"), ("R-1K", "1k"), ("R-100", "100R")):
        create(client, "Resistors", res(pn, value))
    return client


def test_create_get_update_awkward_part_numbers(client: TestClient) -> None:
    fuse = {"Manufacturer": "Bourns", "Protection Type": "PTC (rückstellend)", "Current Rating": "3A",
            "Library Ref": "FUSE", "Footprint Ref": "FUSC3225X100N"}
    for pn in ("MF-GSMF300/36X-2", "LY Q396", "2N7002,215", "LTC3128EUFD#PBF", "A%B"):
        out = create(client, "Fuses", {**fuse, "Part Number": pn, "Manufacturer Part Number": pn})
        assert out["part"]["Part Number"] == pn and out["warnings"] == []
        got = client.get(url("Fuses", pn))
        assert got.status_code == 200 and got.json()["Part Number"] == pn
        patched = client.patch(url("Fuses", pn), json={"values": {"Voltage Rating": "60"}})
        assert patched.status_code == 200 and patched.json()["part"]["Voltage Rating"] == "60V"
        hist = client.get(url("Fuses", pn, "/history"))
        assert hist.status_code == 200 and [e["action"] for e in hist.json()] == ["update", "create"]
    raw_slash = client.get("/api/parts/Fuses/MF-GSMF300/36X-2")  # unencoded slash works too
    assert raw_slash.status_code == 200
    assert client.get(url("Fuses", "nope")).status_code == 404
    assert client.get(url("Fuses", "nope", "/history")).status_code == 404
    assert client.patch(url("Fuses", "nope"), json={"values": {"Voltage Rating": "1"}}).status_code == 404
    assert client.get(url("Gibtsnicht", "x")).status_code == 404


def test_create_errors_and_warnings(client: TestClient) -> None:
    r = client.post("/api/parts/Resistors", json={"values": {"Manufacturer": "YAGEO", "Manufacturer Part Number": "X"}})
    assert r.status_code == 400
    detail = r.json()["detail"]
    assert "Pflichtfeld" in detail["message"] and "Resistance" in detail["errors"]
    out = create(client, "Resistors", res("W-1", "1k", Size="9999"))
    assert any("Auswahlliste" in w for w in out["warnings"])
    capacitor = {"Part Number": "w-1", "Manufacturer": "Murata", "Manufacturer Part Number": "w-1", "Capacitance": "1n",
                 "Voltage": "50", "Library Ref": "CAP", "Footprint Ref": "CAPC1608X87N"}
    dup = client.post("/api/parts/Capacitors", json={"values": capacitor})
    assert dup.status_code == 409 and "Resistors" in dup.json()["detail"]
    unknown = client.post("/api/parts/Capacitors", json={"values": {**capacitor, "Part Number": "C-9", "Resistance": "1k"}})
    assert unknown.status_code == 400 and "Resistance" in unknown.json()["detail"]["errors"]
    draft = client.post("/api/parts/Resistors", json={"values": {"Part Number": "D-1", "Resistance": "1k"}, "allow_incomplete": True})
    assert draft.status_code == 201 and any("Pflichtfeld" in w for w in draft.json()["warnings"])
    assert client.post("/api/parts/Resistors", json={"wrong": 1}).status_code == 422
    assert client.post("/api/parts/Gibtsnicht", json={"values": {}}).status_code == 404


def test_list_query_and_sort(filled: TestClient) -> None:
    client = filled
    page = client.get("/api/parts/Resistors", params={"search": "4.7k"}).json()
    assert page["total"] == 1 and page["rows"][0]["Part Number"] == "R-4K7" and page["limit"] == 200
    filters = json.dumps([{"column": "Resistance", "op": "gt", "value": "500"}])
    page = client.get("/api/parts/Resistors", params={"filters": filters, "sort": "-Resistance", "limit": 2, "offset": 1}).json()
    assert page["total"] == 3 and [r["Part Number"] for r in page["rows"]] == ["R-4K7", "R-1K"]
    page = client.get("/api/parts/Resistors", params={"sort": json.dumps([{"column": "Resistance", "desc": False}]), "limit": 0}).json()
    assert [r["Part Number"] for r in page["rows"]] == ["R-100", "R-1K", "R-4K7", "R-10K"] and page["limit"] is None
    single = client.get("/api/parts/Resistors", params={"filters": json.dumps({"column": "Size", "op": "eq", "value": "0603"})})
    assert single.json()["total"] == 4
    body = {"search": "yageo", "filters": [{"column": "Resistance", "op": "le", "value": "1k"}], "sort": [{"column": "Resistance"}]}
    posted = client.post("/api/parts/Resistors/query", json=body).json()
    assert [r["Part Number"] for r in posted["rows"]] == ["R-100", "R-1K"]
    assert client.get("/api/parts/Resistors", params={"filters": "{kaputt"}).status_code == 400
    assert client.get("/api/parts/Resistors", params={"filters": json.dumps([{"column": "Nope", "op": "eq", "value": 1}])}).status_code == 400
    assert client.get("/api/parts/Resistors", params={"filters": json.dumps([{"column": "Size", "op": "like"}])}).status_code == 400
    assert client.get("/api/parts/Resistors", params={"sort": "Nope"}).status_code == 400
    assert client.get("/api/parts/Resistors", params={"sort": "[kaputt"}).status_code == 400
    assert client.get("/api/parts/Resistors", params={"limit": 20000}).status_code == 422
    assert client.get("/api/parts/Gibtsnicht").status_code == 404


def test_delete_bulk_duplicate(filled: TestClient) -> None:
    client = filled
    bulk = client.post("/api/parts/Resistors/bulk-update", json={"part_numbers": ["R-1K", "R-100"], "values": {"Tolerance": "5"}})
    assert bulk.status_code == 200 and bulk.json()["updated"] == 2
    assert client.get(url("Resistors", "R-1K")).json()["Tolerance"] == "±5%"
    bad_bulk = client.post("/api/parts/Resistors/bulk-update", json={"part_numbers": ["R-1K"], "values": {"Part Number": "X"}})
    assert bad_bulk.status_code == 400
    dup = client.post(url("Resistors", "R-1K", "/duplicate"), json={"new_part_number": "R-1K-B"})
    assert dup.status_code == 201 and dup.json()["part"]["Part Number"] == "R-1K-B"
    assert client.post(url("Resistors", "R-1K", "/duplicate"), json={"new_part_number": "r-10k"}).status_code == 409
    assert client.post(url("Resistors", "nope", "/duplicate"), json={"new_part_number": "x"}).status_code == 404
    deleted = client.request("DELETE", "/api/parts/Resistors", json={"part_numbers": ["R-1K-B", "nope"]})
    assert deleted.status_code == 200
    assert deleted.json() == {"deleted": 1, "backup": None, "part_numbers": ["R-1K-B"], "matched": None}
    alias = client.post("/api/parts/Resistors/delete", json={"part_numbers": ["R-100"]})
    assert alias.json()["deleted"] == 1
    many = [f"M-{i}" for i in range(21)]
    db_parts.upsert_parts("Resistors", [res(pn, f"{i + 1}k") for i, pn in enumerate(many)])
    big = client.request("DELETE", "/api/parts/Resistors", json={"part_numbers": many})
    assert big.json()["deleted"] == 21 and big.json()["backup"].startswith("PartForge_")


def test_validate_render_smart_fill_regenerate_distinct(filled: TestClient) -> None:
    client = filled
    check = client.post("/api/parts/Resistors/validate", json={"values": res("R-10K", "10k")}).json()
    assert check["conflict"]["part_number"] == "R-10K" and check["rendered"]["Comment"] == "10k"
    check = client.post("/api/parts/Resistors/validate", json={"values": {"Tolerance": "J"}, "part_number": "R-1K"}).json()
    assert check["conflict"] is None and check["warnings"] and check["values"]["Tolerance"] == "J"
    assert client.post("/api/parts/Resistors/validate", json={"values": {}, "part_number": "nope"}).status_code == 404
    rendered = client.post("/api/parts/Resistors/render-templates", json={"values": {"Resistance": "2k2", "Size": "0402"}}).json()
    assert rendered == {"Value": "2.2kΩ", "Comment": "2k2", "Description": "RES 0402 2.2kΩ"}
    fill = client.post("/api/parts/Resistors/smart-fill", json={"text": "10k 1% 0603 Yageo"}).json()
    assert fill["values"]["Resistance"] == "10kΩ" and fill["unmatched"] == ["Yageo"]
    client.patch("/api/categories/Resistors", json={"description_template": "R {Resistance}"})
    regen = client.post("/api/parts/Resistors/regenerate-texts", json={})
    assert regen.json()["updated"] == 4 and regen.json()["backup"] is None  # 4 parts: no automatic backup
    assert len(regen.json()["audit_ids"]) == 4  # undo handle (POST /parts/{table}/restore-updates)
    assert client.get(url("Resistors", "R-1K")).json()["Description"] == "R 1kΩ"
    distinct = client.get("/api/distinct/Resistors", params={"column": "Resistance"}).json()
    assert [d["value"] for d in distinct] == ["100Ω", "1kΩ", "4.7kΩ", "10kΩ"]
    assert client.get("/api/distinct/Resistors", params={"column": "Nope"}).status_code == 404


def test_search_stats_health(filled: TestClient) -> None:
    client = filled
    hits = client.get("/api/search", params={"q": "r-4k7"}).json()
    assert hits[0]["part_number"] == "R-4K7" and hits[0]["label"] == "Widerstände"
    assert client.get("/api/search", params={"q": ""}).json() == []
    stats = client.get("/api/stats").json()
    assert stats["total_parts"] == 4 and stats["missing"]["datasheet"] == 4 and stats["recent_changes"]
    assert stats["symbols"]["total"] == 0 and stats["status_counts"] == {"Entwurf": 4}
    report = client.get("/api/health-check").json()
    assert report["summary"]["symbol_missing"] == 4 and report["summary"]["datasheet_missing"] == 4
    assert report["issues"][0]["severity"] in ("error", "warning")
    limited = client.get("/api/health-check", params={"max_per_check": 1}).json()
    assert limited["truncated"] and len(limited["issues"]) == len(limited["summary"])


def test_import_preview_and_job(client: TestClient) -> None:
    csv_text = ("Teilenummer;Hersteller;Widerstand;Toleranz;Leistung;Symbol;Footprint\r\n"
                "RC0603FR-0710KL;YAGEO;10k;1%;0,1W;RES;RESC1608X55N\r\n"
                "RC0603FR-074K7L;YAGEO;4k7;1%;0,1W;RES;RESC1608X55N\r\n")
    files = {"file": ("Liste.csv", csv_text.encode("cp1252"), "text/csv")}
    r = client.post("/api/import/preview", files=files, data={"table": "Resistors"})
    assert r.status_code == 200, r.text
    prev = r.json()
    assert prev["file_name"] == "Liste.csv" and prev["row_count"] == 2 and prev["delimiter"] == ";"
    assert {m["header"]: m["column"] for m in prev["mapping"]}["Widerstand"] == "Resistance"
    assert [s["status"] for s in prev["samples"]] == ["neu", "neu"]
    again = client.post(f"/api/import/preview/{prev['token']}", json={"table": "Capacitors", "header_row": 0})
    assert again.status_code == 200 and again.json()["header_row"] is None and again.json()["row_count"] == 3
    # the file has no MPN column (required): import the rows as drafts
    run = client.post("/api/import/run", json={"token": prev["token"], "table": "resistors", "mapping": prev["mapping"],
                                               "mode": "skip", "allow_incomplete": True})
    assert run.status_code == 200, run.text
    job = wait_for_job(client, run.json()["job"]["id"])
    assert job["status"] == "done", job
    assert job["result"]["created"] == 2 and job["result"]["backup"].startswith("PartForge_")
    assert client.get(url("Resistors", "RC0603FR-074K7L")).json()["Resistance"] == "4.7kΩ"
    bad_upload = client.post("/api/import/preview", files={"file": ("x.pdf", b"%PDF", "application/pdf")})
    assert bad_upload.status_code == 400 and "nicht unterstützt" in bad_upload.json()["detail"]
    assert client.post("/api/import/preview", files=files, data={"header_row": "zwei"}).status_code == 400
    assert client.post("/api/import/run", json={"token": "0" * 24, "table": "Resistors", "mapping": []}).status_code == 400
    assert client.post("/api/import/run", json={"token": prev["token"], "table": "Nope", "mapping": []}).status_code == 404
    assert client.post("/api/import/preview/zzz", json={}).status_code == 400


def test_export_endpoints(filled: TestClient) -> None:
    client = filled
    r = client.get("/api/export/Resistors", params={"format": "csv", "pns": ["R-1K", "R-10K"], "columns": ["Part Number", "Resistance"]})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"] and "Resistors_" in r.headers["content-disposition"]
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "Part Number;Resistance" and sorted(lines[1:]) == ["R-10K;10kΩ", "R-1K;1kΩ"]
    x = client.get("/api/export/Resistors", params={"format": "xlsx", "search": "4k7"})
    assert x.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert sheet.max_row == 2 and sheet["A2"].value == "R-4K7"
    filtered = client.get("/api/export/Resistors", params={"format": "csv", "filters": json.dumps([{"column": "Resistance", "op": "lt", "value": "1k"}])})
    assert filtered.content.decode("utf-8-sig").splitlines()[1].startswith("R-100;")
    assert client.get("/api/export/Resistors", params={"format": "pdf"}).status_code == 422
    assert client.get("/api/export/Resistors", params={"columns": ["Nope"]}).status_code == 400
    saved = client.post("/api/export/Resistors", json={"format": "xlsx", "part_numbers": ["R-1K"]})
    assert saved.status_code == 200 and saved.json()["rows"] == 1 and saved.json()["rel_path"].startswith("Exports\\")


def test_backup_endpoints(filled: TestClient) -> None:
    client = filled
    made = client.post("/api/backups", json={"reason": "vor Test"})
    assert made.status_code == 201 and made.json()["reason"] == "vor_Test"
    assert client.post("/api/backups").status_code == 201  # body optional
    listing = client.get("/api/backups").json()
    assert len(listing) == 2 and listing[0]["rel_path"].startswith("Backups\\")
    client.request("DELETE", "/api/parts/Resistors", json={"part_numbers": ["R-1K"]})
    restored = client.post(f"/api/backups/{made.json()['name']}/restore")
    assert restored.status_code == 200
    body = restored.json()
    assert body["ok"] and body["restored"] == made.json()["name"]
    # the library files are compared with the restored registry (INT-06)
    assert body["job"]["kind"] == "libraries.scan"
    assert wait_for_job(client, body["job"]["id"], timeout=60)["status"] == "done"
    assert client.get(url("Resistors", "R-1K")).status_code == 200
    assert client.post("/api/backups/PartForge_20000101-000000_nope.sqlite/restore").status_code == 404


def test_saved_view_endpoints(client: TestClient) -> None:
    created = client.post("/api/views", json={"name": "Kleine Werte", "table": "Resistors",
                                              "filters": [{"column": "Resistance", "op": "lt", "value": "1k"}],
                                              "sort": [{"column": "Resistance", "desc": True}], "columns": ["Part Number"]})
    assert created.status_code == 201, created.text
    view = created.json()
    assert view["id"] and view["table"] == "Resistors" and view["author"]
    assert [v["name"] for v in client.get("/api/views", params={"table": "Resistors"}).json()] == ["Kleine Werte"]
    assert client.get("/api/views", params={"table": "Capacitors"}).json() == []
    assert client.get(f"/api/views/{view['id']}").json()["sort"][0]["desc"] is True
    patched = client.patch(f"/api/views/{view['id']}", json={"name": "Umbenannt"})
    assert patched.status_code == 200 and patched.json()["name"] == "Umbenannt" and patched.json()["filters"] == view["filters"]
    assert client.patch(f"/api/views/{view['id']}", json={"columns": ["Nope"]}).status_code == 400
    assert client.post("/api/views", json={"name": "x", "table": "Nope"}).status_code == 404
    assert client.delete(f"/api/views/{view['id']}").json() == {"deleted": True}
    assert client.delete(f"/api/views/{view['id']}").status_code == 404
    assert client.get(f"/api/views/{view['id']}").status_code == 404


def test_new_library_with_examples(bare_client: TestClient, tmp_path: Any) -> None:
    from partforge.db import examples

    created = bare_client.post("/api/workspace/create", json={"path": str(tmp_path / "Lib"), "name": "Demo",
                                                              "with_examples": True, "write_libraries": False})
    assert created.status_code == 200, created.text
    try:
        parts_count = len(examples.EXAMPLE_PARTS)  # 19 parts, 15 symbols, 13 footprints (ARCHITEKTUR §4)
        assert created.json()["examples"] == {"parts": parts_count, "symbols": 15, "footprints": 13}
        assert created.json()["job"] is None  # write_libraries=False: the symbols/footprints stay pending
        stats = bare_client.get("/api/stats").json()
        assert stats["total_parts"] == parts_count
        assert stats["symbols"]["by_status"]["pending"] == 15 and stats["footprints"]["by_status"]["pending"] == 13
        hit = bare_client.get("/api/search", params={"q": "LM358"}).json()[0]
        assert hit["part_number"] == "LM358DR" and hit["table"] == "Amplifiers"
        page = bare_client.get("/api/parts/Resistors", params={"sort": "Resistance"}).json()
        assert [r["Comment"] for r in page["rows"]] == ["100R", "1k", "4k7", "10k"]
        caps = bare_client.get("/api/parts/Capacitors", params={"sort": "Capacitance"}).json()
        assert [r["Comment"] for r in caps["rows"]] == ["100n 50V", "10u 25V"]  # the Comment carries the voltage
    finally:
        assert bare_client.post("/api/workspace/close").status_code == 200


def test_parts_routes_require_open_library(bare_client: TestClient) -> None:
    for method, path in (("get", "/api/parts/Resistors"), ("get", "/api/search?q=x"), ("get", "/api/stats"),
                         ("get", "/api/health-check"), ("get", "/api/backups"), ("get", "/api/views"),
                         ("get", "/api/export/Resistors")):
        r = getattr(bare_client, method)(path)
        assert r.status_code == 409, (path, r.status_code)


# --------------------------------------------------------------------------- regressions (review)


def test_locked_database_gives_german_409(client: TestClient, workspace: Any) -> None:
    """Regression: another program holding the write lock produced an English HTTP 500."""
    import sqlite3

    blocker = sqlite3.connect(str(workspace.sqlite_path), timeout=1, isolation_level=None)
    try:
        blocker.execute("BEGIN IMMEDIATE")  # e.g. a long write of another program
        r = client.post("/api/parts/Resistors", json={"values": res("R-LOCK", "1k")})
        assert r.status_code == 409, r.text
        assert "gesperrt" in r.json()["detail"] and "erneut versuchen" in r.json()["detail"]
    finally:
        blocker.execute("ROLLBACK")
        blocker.close()
    assert client.post("/api/parts/Resistors", json={"values": res("R-LOCK", "1k")}).status_code == 201


def test_unexpected_database_errors_are_german(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    import sqlite3

    from sqlalchemy.exc import OperationalError

    from partforge.db import categories as db_categories

    def locked(*_args: Any, **_kwargs: Any) -> None:
        raise OperationalError("ALTER TABLE", {}, sqlite3.OperationalError("database is locked"))

    def broken(*_args: Any, **_kwargs: Any) -> None:
        raise OperationalError("SELECT", {}, sqlite3.OperationalError("disk I/O error"))

    monkeypatch.setattr(db_categories, "update_category", locked)
    r = client.patch("/api/categories/Resistors", json={"label": "X"})
    assert r.status_code == 409 and "gesperrt" in r.json()["detail"]
    monkeypatch.setattr(db_parts, "stats", broken)
    r = client.get("/api/stats")
    assert r.status_code == 500 and r.json()["detail"].startswith("Unerwarteter Datenbankfehler: disk I/O error")


def test_unicode_search_via_api(client: TestClient) -> None:
    create(client, "Resistors", res("R-UML", "10k", Notes="Übertemperatur geprüft"))
    page = client.get("/api/parts/Resistors", params={"search": "ÜBERTEMPERATUR"}).json()
    assert page["total"] == 1 and page["rows"][0]["Part Number"] == "R-UML"
    filters = json.dumps([{"column": "Notes", "op": "contains", "value": "übertemp"}])
    assert client.get("/api/parts/Resistors", params={"filters": filters}).json()["total"] == 1
    assert client.get("/api/parts/Resistors", params={"search": "kΩ"}).json()["total"] == 1


def test_import_preview_reports_header_units(client: TestClient) -> None:
    csv_text = "MPN;Hersteller;Kapazität (pF);Spannung (V);Symbol;Footprint\r\nGRM1885C1H101JA01D;Murata;100;50;CAP;CAPC1608X87N\r\n"
    files = {"file": ("caps.csv", csv_text.encode("cp1252"), "text/csv")}
    prev = client.post("/api/import/preview", files=files, data={"table": "Capacitors"}).json()
    item = next(m for m in prev["mapping"] if m["header"] == "Kapazität (pF)")
    assert item["column"] == "Capacitance" and item["unit_hint"] == "pF" and "pF" in item["unit_note"]
    assert prev["samples"][0]["values"]["Capacitance"] == "100pF"
    run = client.post("/api/import/run", json={"token": prev["token"], "table": "Capacitors", "mapping": prev["mapping"]})
    job = wait_for_job(client, run.json()["job"]["id"])
    assert job["status"] == "done" and job["result"]["created"] == 1, job
    assert client.get(url("Capacitors", "GRM1885C1H101JA01D")).json()["Capacitance"] == "100pF"
    bad = [{**m, "unit_hint": "V"} if m["header"] == "Kapazität (pF)" else m for m in prev["mapping"]]
    run = client.post("/api/import/run", json={"token": prev["token"], "table": "Capacitors", "mapping": bad})
    job = wait_for_job(client, run.json()["job"]["id"])
    assert job["status"] == "failed" and "passt nicht" in (job["error"] or "")


def test_excel_safe_csv_export_and_literal_xlsx(filled: TestClient) -> None:
    client = filled
    client.patch(url("Resistors", "R-1K"), json={"values": {"Comment": "=Value"}})
    safe = client.get("/api/export/Resistors", params={"format": "csv", "pns": ["R-1K"], "columns": ["Part Number", "Comment"]})
    assert safe.content.decode("utf-8-sig").splitlines()[1] == "R-1K;'=Value"  # guarded by default (INT-09)
    plain = client.get("/api/export/Resistors", params={"format": "csv", "pns": ["R-1K"], "columns": ["Part Number", "Comment"],
                                                        "excel_safe": False})
    assert plain.content.decode("utf-8-sig").splitlines()[1] == "R-1K;=Value"
    x = client.get("/api/export/Resistors", params={"format": "xlsx", "pns": ["R-1K"], "columns": ["Part Number", "Comment"]})
    sheet = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert sheet["B2"].value == "=Value" and sheet["B2"].data_type == "s"
    saved = client.post("/api/export/Resistors", json={"format": "csv", "part_numbers": ["R-1K"], "excel_safe": True})
    assert saved.status_code == 200 and saved.json()["rows"] == 1


def test_export_as_job(filled: TestClient) -> None:
    client = filled
    started = client.post("/api/export/Resistors/job", json={"format": "xlsx", "columns": ["Part Number", "Resistance"]})
    assert started.status_code == 200, started.text
    job = wait_for_job(client, started.json()["job"]["id"])
    assert job["status"] == "done", job
    result = job["result"]
    assert result["rows"] == 4 and result["rel_path"].startswith("Exports\\") and result["file_name"].endswith(".xlsx")
    sheet = openpyxl.load_workbook(result["path"]).active
    assert sheet.max_row == 5 and sheet["B1"].value == "Resistance"
    assert client.post("/api/export/Resistors/job", json={"columns": ["Nope"]}).status_code == 400
    assert client.post("/api/export/Nope/job", json={}).status_code == 404


def test_backups_work_while_database_is_closed(filled: TestClient) -> None:
    from partforge.db import close_database, core

    client = filled
    made = client.post("/api/backups", json={"reason": "stand"}).json()
    client.request("DELETE", "/api/parts/Resistors", json={"part_numbers": ["R-1K"]})
    close_database()  # e.g. after a failed open; the workspace stays the current one
    assert client.get("/api/parts/Resistors").status_code == 409
    assert [b["name"] for b in client.get("/api/backups").json()] == [made["name"]]
    restored = client.post(f"/api/backups/{made['name']}/restore")
    assert restored.status_code == 200 and core.is_open()
    assert client.get(url("Resistors", "R-1K")).status_code == 200
