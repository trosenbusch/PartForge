"""Spreadsheet import (CSV/XLSX/ODS, mapping, preview, job) and export (CSV/XLSX round trip)."""
from __future__ import annotations

import io
import os
import time
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from partforge.db import backup, categories, parts
from partforge.db.examples import seed_examples
from partforge.importers import spreadsheet as ss
from partforge.importers.spreadsheet import ImportOptions, MappingItem, SpreadsheetError
from partforge.jobs import JobManager
from partforge.workspace import Workspace

GERMAN_CSV = (
    "Teilenummer;Hersteller;Hersteller-Teilenummer;Widerstand;Toleranz;Leistung;Baugröße;Symbol;Footprint;Bemerkung\r\n"
    "RC0603FR-0710KL;YAGEO;RC0603FR-0710KL;10k;1%;0,1W;0603;RES;RESC1608X55N;Prüfung ok (µ-Test)\r\n"
    "RC0603FR-074K7L;YAGEO;RC0603FR-074K7L;4,7k;1 %;1/10W;603;RES;RESC1608X55N;\r\n"
    ";;;;;;;;;\r\n"
    "RC0603FR-071KL;YAGEO;RC0603FR-071KL;1k;1;0,1;0603;RES;RESC1608X55N;\"mit; Semikolon\"\r\n"
)


def upload(content: bytes, name: str) -> tuple[str, Path]:
    token = ss.store_upload(content, name)
    path, original = ss.upload_path(token)
    assert original == name
    return token, path


def mapping_of(preview: dict[str, Any]) -> dict[str, str | None]:
    return {m["header"]: m["column"] for m in preview["mapping"]}


def test_german_excel_csv_preview_and_import(workspace: Workspace) -> None:
    token, path = upload(GERMAN_CSV.encode("cp1252"), "Widerstände.csv")
    prev = ss.preview(path, table="Resistors")
    assert prev["encoding"] == "cp1252" and prev["delimiter"] == ";" and prev["header_row"] == 1
    assert prev["row_count"] == 3 and prev["rows"][0][9] == "Prüfung ok (µ-Test)"
    assert mapping_of(prev) == {
        "Teilenummer": "Part Number", "Hersteller": "Manufacturer", "Hersteller-Teilenummer": "Manufacturer Part Number",
        "Widerstand": "Resistance", "Toleranz": "Tolerance", "Leistung": "Power", "Baugröße": "Size",
        "Symbol": "Library Ref", "Footprint": "Footprint Ref", "Bemerkung": "Notes",
    }
    assert [s["status"] for s in prev["samples"]] == ["neu", "neu", "neu"]
    assert prev["samples"][1]["values"]["Resistance"] == "4,7k"
    before = len(backup.list_backups())
    result = ss.run_import(None, ImportOptions(token=token, table="Resistors", mapping=prev["mapping"], author="Import"))
    assert result["created"] == 3 and result["error_count"] == 0 and result["rows"] == 3
    assert result["backup"] and len(backup.list_backups()) == before + 1
    row = parts.get_part("Resistors", "RC0603FR-074K7L")
    assert row and row["Resistance"] == "4.7kΩ" and row["Power"] == "0.1W" and row["Size"] == "0603"
    assert row["PF Author"] == "Import" and row["Description"] == "RES 0603 4.7kΩ ±1% 0.1W"
    assert parts.get_part("Resistors", "RC0603FR-071KL")["Notes"] == "mit; Semikolon"  # type: ignore[index]
    assert parts.get_part("Resistors", "RC0603FR-0710KL")["Notes"] == "Prüfung ok (µ-Test)"  # type: ignore[index]
    again = ss.preview(path, table="Resistors")
    assert [s["status"] for s in again["samples"]] == ["vorhanden", "vorhanden", "vorhanden"]


def test_english_utf8_csv_with_commas_and_defaults(workspace: Workspace) -> None:
    text = ('﻿Mfr Part Number,Capacitance (F),Voltage Rating,Tolerance,Datasheet,Description\n'
            'GRM188R71H104KA93D,100n,50V,±10%,www.murata.com/a.pdf,"CAP, 100nF X7R"\n'
            'GRM188R71H103KA01D,10n,50V,±10%,,\n')
    token, path = upload(text.encode("utf-8"), "caps.csv")
    prev = ss.preview(path, table="Capacitors")
    assert prev["encoding"] == "utf-8-sig" and prev["delimiter"] == ","
    assert mapping_of(prev) == {"Mfr Part Number": "Manufacturer Part Number", "Capacitance (F)": "Capacitance",
                                "Voltage Rating": "Voltage", "Tolerance": "Tolerance", "Datasheet": "ComponentLink1URL",
                                "Description": "Description"}
    result = ss.run_import(None, {"token": token, "table": "capacitors", "mapping": prev["mapping"],
                                  "defaults": {"Manufacturer": "Murata", "Library Ref": "CAP",
                                               "Footprint Ref": "CAPC1608X87N", "Dielectric": None}})
    assert result["created"] == 2, result
    first = parts.get_part("Capacitors", "GRM188R71H104KA93D")
    assert first and first["Description"] == "CAP, 100nF X7R" and first["ComponentLink1URL"] == "https://www.murata.com/a.pdf"
    assert first["ComponentLink1Description"] == "Datenblatt" and first["Manufacturer"] == "Murata"
    second = parts.get_part("Capacitors", "GRM188R71H103KA01D")
    assert second and second["Description"] == "CAP 10nF 50V ±10%"  # generated (empty cell)


def test_xlsx_and_update_mode(workspace: Workspace, tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kondensatoren"
    other = wb.create_sheet("Leer")
    other["A1"] = "nichts"
    ws.append(["Part Number", "Manufacturer", "Manufacturer Part Number", "Capacitance", "Voltage", "Size", "Library Ref",
               "Footprint Ref", "Polarized"])
    ws.append(["C-1", "Murata", "C-1", 1e-07, 50, 603, "CAP", "CAPC1608X87N", True])
    ws.append(["C-2", "Murata", "C-2", 4.7e-06, 6.3, "0805", "CAP", "CAPC2012X145N", False])
    buffer = io.BytesIO()
    wb.save(buffer)
    token, path = upload(buffer.getvalue(), "liste.xlsx")
    prev = ss.preview(path, table="Capacitors")
    assert prev["sheets"] == ["Kondensatoren", "Leer"] and prev["sheet"] == "Kondensatoren"
    assert all(m["score"] == 100.0 for m in prev["mapping"])
    assert prev["rows"][0][3] == "0.0000001"
    result = ss.run_import(None, ImportOptions(token=token, table="Capacitors", mapping=prev["mapping"]))
    assert result["created"] == 2
    row = parts.get_part("Capacitors", "C-1")
    assert row and row["Capacitance"] == "100nF" and row["Voltage"] == "50V" and row["Size"] == "0603" and row["Polarized"] == "Ja"
    assert parts.get_part("Capacitors", "C-2")["Voltage"] == "6.3V"  # type: ignore[index]
    # second import of the same file: update mode -> unchanged; skip mode -> skipped
    again = ss.run_import(None, ImportOptions(token=token, table="Capacitors", mapping=prev["mapping"], mode="update"))
    assert again["unchanged"] == 2 and again["updated"] == 0
    skipped = ss.run_import(None, ImportOptions(token=token, table="Capacitors", mapping=prev["mapping"]))
    assert skipped["skipped"] == 2
    empty_sheet = ss.preview(path, table="Capacitors", sheet="Leer")
    assert empty_sheet["row_count"] == 0 and empty_sheet["headers"] == ["nichts"]


def test_ods_import(workspace: Workspace, tmp_path: Path) -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name="Teile")

    def row(*values: Any, repeat_empty: int = 0) -> TableRow:
        tr = TableRow()
        for value in values:
            if isinstance(value, (int, float)):
                cell = TableCell(valuetype="float", value=str(value))
                cell.addElement(P(text=str(value)))
            else:
                cell = TableCell(valuetype="string")
                cell.addElement(P(text=value))
            tr.addElement(cell)
        if repeat_empty:
            tr.addElement(TableCell(numbercolumnsrepeated=str(repeat_empty)))
        return tr

    table.addElement(row("Teilenummer", "Hersteller", "Impedanz", "Nennstrom", "Symbol", "Footprint", repeat_empty=1000))
    table.addElement(row("BLM18PG221SN1D", "Murata", 220, 1.4, "FERRITE", "INDC1608X95N"))
    table.addElement(TableRow(numberrowsrepeated="1048000"))
    doc.spreadsheet.addElement(table)
    path = tmp_path / "ferrite.ods"
    doc.save(str(path))
    token, stored = upload(path.read_bytes(), "ferrite.ods")
    prev = ss.preview(stored, table="Ferrite Beads")
    assert prev["sheets"] == ["Teile"] and prev["row_count"] == 1 and len(prev["headers"]) == 6
    assert mapping_of(prev)["Impedanz"] == "Impedance" and mapping_of(prev)["Nennstrom"] == "Current Rating"
    result = ss.run_import(None, ImportOptions(token=token, table="Ferrite Beads", mapping=prev["mapping"],
                                               defaults={"Manufacturer Part Number": "BLM18PG221SN1D"}))
    assert result["created"] == 1
    row_ = parts.get_part("Ferrite Beads", "BLM18PG221SN1D")
    assert row_ and row_["Impedance"] == "220Ω" and row_["Current Rating"] == "1.4A"


def test_header_detection_and_explicit_header_row(workspace: Workspace) -> None:
    no_header = "1k;0603;1%\n2k;0603;1%\n"
    _, path = upload(no_header.encode(), "ohne.csv")
    prev = ss.preview(path, table="Resistors")
    assert prev["header_row"] is None and prev["headers"] == ["Spalte A", "Spalte B", "Spalte C"] and prev["row_count"] == 2
    titled = "Bestellliste Oktober;;\nTeilenummer;Widerstand;Toleranz\nR-1;1k;1%\n"
    _, path2 = upload(titled.encode(), "titel.csv")
    prev2 = ss.preview(path2, table="Resistors", header_row=2)
    assert prev2["headers"] == ["Teilenummer", "Widerstand", "Toleranz"] and prev2["row_count"] == 1
    prev3 = ss.preview(path2, header_row=None)
    assert prev3["header_row"] is None and prev3["row_count"] == 3 and prev3["mapping"] == []
    with pytest.raises(SpreadsheetError, match="Kopfzeile"):
        ss.preview(path2, header_row=99)
    tsv = "Teilenummer\tWiderstand\nR-1\t1k\n"
    _, path4 = upload(tsv.encode(), "liste.tsv")
    assert ss.preview(path4)["delimiter"] == "\t"


def test_mapping_validation_errors(workspace: Workspace) -> None:
    token, _ = upload(GERMAN_CSV.encode("cp1252"), "w.csv")
    base = dict(token=token, table="Resistors")
    with pytest.raises(SpreadsheetError, match="mehreren Spalten"):
        ss.run_import(None, ImportOptions(**base, mapping=[MappingItem(index=0, column="Part Number"),
                                                          MappingItem(index=2, column="part number")]))
    with pytest.raises(SpreadsheetError, match="gibt es in der Kategorie"):
        ss.run_import(None, ImportOptions(**base, mapping=[MappingItem(index=0, column="Gibtsnicht")]))
    with pytest.raises(SpreadsheetError, match="Spalte 100"):
        ss.run_import(None, ImportOptions(**base, mapping=[MappingItem(index=99, column="Part Number")]))
    with pytest.raises(SpreadsheetError, match="Keine Spalte"):
        ss.run_import(None, ImportOptions(**base, mapping=[MappingItem(index=0, column=None)]))
    with pytest.raises(KeyError):
        ss.run_import(None, ImportOptions(token=token, table="Gibtsnicht", mapping=[]))


def test_upload_validation(workspace: Workspace, monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(SpreadsheetError, match="nicht unterstützt"):
        ss.store_upload(b"x", "liste.pdf")
    with pytest.raises(SpreadsheetError, match="leer"):
        ss.store_upload(b"", "liste.csv")
    monkeypatch.setattr(ss, "MAX_UPLOAD_BYTES", 10)
    with pytest.raises(SpreadsheetError, match="zu groß"):
        ss.store_upload(b"x" * 11, "liste.csv")
    for bad in ("../../etc", "zzzzzzzzzzzzzzzzzzzzzzzz", ""):
        with pytest.raises(SpreadsheetError):
            ss.upload_path(bad)
    monkeypatch.setattr(ss, "MAX_UPLOAD_BYTES", 1024)
    token = ss.store_upload(b"a;b\n1;2\n", "alt.csv")
    old, _ = ss.upload_path(token)
    stamp = time.time() - ss.UPLOAD_MAX_AGE_S - 10
    os.utime(old, (stamp, stamp))
    ss.store_upload(b"a;b\n1;2\n", "neu.csv")  # triggers the cleanup
    assert not old.exists()
    with pytest.raises(SpreadsheetError):
        ss.read_table_file(Path("x.doc"))


def test_import_as_job(workspace: Workspace) -> None:
    token, path = upload(GERMAN_CSV.encode("cp1252"), "job.csv")
    prev = ss.preview(path, table="Resistors")
    manager = JobManager(max_workers=1)
    job = manager.run_sync("import", "Import", ss.run_import, ImportOptions(token=token, table="Resistors", mapping=prev["mapping"]))
    assert job.status == "done" and job.result["created"] == 3
    assert any("3 neu" in entry["text"] for entry in job.log)
    failing = manager.run_sync("import", "Import", ss.run_import, ImportOptions(token="0" * 24, table="Resistors", mapping=[]))
    assert failing.status == "failed" and "nicht mehr vorhanden" in (failing.error or "")


def test_export_csv_roundtrip(workspace: Workspace) -> None:
    seed_examples(workspace)
    original = {r["Part Number"]: r for r in parts.query_parts("Resistors", limit=None)[0]}
    content, name, media = ss.export_parts("Resistors", fmt="csv")
    assert name.startswith("Resistors_") and name.endswith(".csv") and media.startswith("text/csv")
    assert content.startswith(b"\xef\xbb\xbf") and b"\r\n" in content
    header = content.decode("utf-8-sig").split("\r\n", 1)[0]
    assert header.split(";")[:3] == ["Part Number", "Manufacturer", "Manufacturer Part Number"]
    token, path = upload(content, name)
    prev = ss.preview(path, table="Resistors")
    assert all(m["column"] == m["header"] for m in prev["mapping"])
    assert parts.delete_parts("Resistors", list(original), backup=False) == len(original)
    result = ss.run_import(None, ImportOptions(token=token, table="Resistors", mapping=prev["mapping"]))
    assert result["created"] == len(original) and result["error_count"] == 0
    for pn, row in original.items():
        new = parts.get_part("Resistors", pn)
        assert new is not None
        assert {k: v for k, v in new.items() if not k.startswith("PF ")} == {k: v for k, v in row.items() if not k.startswith("PF ")}
        assert new["PF Status"] == row["PF Status"]


def test_export_xlsx_formatting_and_selection(workspace: Workspace) -> None:
    seed_examples(workspace)
    content, name, media = ss.export_parts("Resistors", fmt="xlsx", pns=["rc0603fr-0710kl", "RC0603FR-071KL"],
                                           columns=["Part Number", "resistance", "Description", "Part Number"])
    assert name.endswith(".xlsx") and "spreadsheetml" in media
    wb = openpyxl.load_workbook(io.BytesIO(content))
    sheet = wb.active
    assert sheet.title == "Resistors"
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Part Number", "Resistance", "Description")
    assert sorted(r[0] for r in rows[1:]) == ["RC0603FR-0710KL", "RC0603FR-071KL"]
    assert sheet["A1"].font.b and sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:C3"
    assert 8 <= sheet.column_dimensions["C"].width <= 60 and sheet.column_dimensions["A"].width >= len("RC0603FR-0710KL")
    filtered = ss.render_export("Resistors", fmt="csv", search="4k7", sort="-Resistance", delimiter=",")
    assert filtered.rows == 1 and filtered.content.decode("utf-8-sig").splitlines()[1].startswith("RC0603FR-074K7L,")
    with pytest.raises(SpreadsheetError, match="Feld"):
        ss.export_parts("Resistors", columns=["Nope"])
    with pytest.raises(SpreadsheetError, match="Exportformat"):
        ss.export_parts("Resistors", fmt="pdf")  # type: ignore[arg-type]
    with pytest.raises(SpreadsheetError, match="Trennzeichen"):
        ss.export_parts("Resistors", fmt="csv", delimiter="x")


def test_export_to_file(workspace: Workspace) -> None:
    seed_examples(workspace)
    path, count = ss.export_to_file("Amplifiers", fmt="xlsx")
    assert path.parent == workspace.exports_dir and path.is_file() and count == 2
    second, _ = ss.export_to_file("Amplifiers", fmt="xlsx")
    assert second != path and second.is_file()
    empty_path, empty_count = ss.export_to_file("Mechanical", fmt="csv")
    assert empty_count == 0 and empty_path.read_text(encoding="utf-8-sig").startswith("Part Number;")
    assert categories.get_category("Mechanical").table == "Mechanical"


# --------------------------------------------------------------------------- regressions (review)


def test_header_units_scale_plain_numbers(workspace: Workspace) -> None:
    """Regression: "Kapazität (pF)" with 100 was stored as 100F (off by 10^12, no warning)."""
    text = ("MPN;Hersteller;Kapazität (pF);Spannung (V);Toleranz (%);Symbol;Footprint\r\n"
            "GRM1885C1H101JA01D;Murata;100;50;5;CAP;CAPC1608X87N\r\n"
            "GRM1885C1H470JA01D;Murata;47p;50;5;CAP;CAPC1608X87N\r\n")  # value with its own unit stays
    token, path = upload(text.encode("cp1252"), "caps.csv")
    prev = ss.preview(path, table="Capacitors")
    hints = {m["header"]: m["unit_hint"] for m in prev["mapping"]}
    assert hints["Kapazität (pF)"] == "pF" and hints["Spannung (V)"] == "V" and hints["MPN"] is None
    note = next(m["unit_note"] for m in prev["mapping"] if m["header"] == "Kapazität (pF)")
    assert note and "pF" in note
    assert prev["samples"][0]["values"]["Capacitance"] == "100pF"  # the preview shows what will be imported
    result = ss.run_import(None, ImportOptions(token=token, table="Capacitors", mapping=prev["mapping"]))
    assert result["created"] == 2 and result["error_count"] == 0 and not result["warnings"]
    assert parts.get_part("Capacitors", "GRM1885C1H101JA01D")["Capacitance"] == "100pF"  # type: ignore[index]
    assert parts.get_part("Capacitors", "GRM1885C1H470JA01D")["Capacitance"] == "47pF"  # type: ignore[index]
    assert parts.get_part("Capacitors", "GRM1885C1H101JA01D")["Tolerance"] == "±5%"  # type: ignore[index]

    res_text = ("MPN;Widerstand [kOhm];Leistung (mW);Toleranz;Symbol;Footprint\r\n"
                "X1;4,7;100;1%;RES;RESC1608X55N\r\nX2;10k;62,5;1%;RES;RESC1608X55N\r\n")
    token2, path2 = upload(res_text.encode("cp1252"), "res.csv")
    prev2 = ss.preview(path2, table="Resistors")
    # a client that leaves out "unit_hint" gets the hint derived from the header again
    lean = [{"index": m["index"], "column": m["column"]} for m in prev2["mapping"]]
    out = ss.run_import(None, ImportOptions(token=token2, table="Resistors", mapping=lean, allow_incomplete=True))
    assert out["created"] == 2, out
    x1 = parts.get_part("Resistors", "X1")
    assert x1 and x1["Resistance"] == "4.7kΩ" and x1["Power"] == "0.1W"
    x2 = parts.get_part("Resistors", "X2")
    assert x2 and x2["Resistance"] == "10kΩ" and x2["Power"] == "62.5mW"


def test_header_unit_problems_and_explicit_hints(workspace: Workspace) -> None:
    cap = categories.get_category("Capacitors")
    capacitance = cap.field("Capacitance")
    manufacturer = cap.field("Manufacturer")
    assert capacitance is not None and manufacturer is not None
    hint, note = ss.header_unit("Kapazität (V)", capacitance)
    assert hint is None and note and "passt nicht" in note
    assert ss.header_unit("Kapazität (typ)", capacitance) == (None, None)  # not a unit
    assert ss.header_unit("Kapazität (@100MHz)", capacitance) == (None, None)
    assert ss.header_unit("Kapazität [µF]", capacitance)[0] == "µF"
    assert ss.header_unit("Hersteller (DE)", manufacturer) == (None, None)
    text = "MPN;Kapazität (V);Symbol;Footprint\r\nC-1;100;CAP;CAPC1608X87N\r\n"
    token, path = upload(text.encode(), "falsch.csv")
    prev = ss.preview(path, table="Capacitors")
    item = next(m for m in prev["mapping"] if m["header"] == "Kapazität (V)")
    assert item["column"] == "Capacitance" and item["unit_hint"] is None and "passt nicht" in item["unit_note"]
    base = [MappingItem(index=0, column="Manufacturer Part Number"), MappingItem(index=2, column="Library Ref"),
            MappingItem(index=3, column="Footprint Ref")]
    with pytest.raises(SpreadsheetError, match="passt nicht"):
        ss.run_import(None, ImportOptions(token=token, table="Capacitors",
                                          mapping=[*base, MappingItem(index=1, column="Capacitance", unit_hint="V")]))
    with pytest.raises(SpreadsheetError, match="keine Einheit"):
        ss.run_import(None, ImportOptions(token=token, table="Capacitors",
                                          mapping=[*base, MappingItem(index=1, column="Notes", unit_hint="nF")]))
    out = ss.run_import(None, ImportOptions(token=token, table="Capacitors", allow_incomplete=True,
                                            mapping=[*base, MappingItem(index=1, column="Capacitance", unit_hint="nF")]))
    assert out["created"] == 1 and parts.get_part("Capacitors", "C-1")["Capacitance"] == "100nF"  # type: ignore[index]


def test_apply_unit_hint() -> None:
    assert ss.apply_unit_hint(100, "pF") == "100pF"
    assert ss.apply_unit_hint(4.7, "kOhm") == "4.7kOhm"
    assert ss.apply_unit_hint("4,7", "kOhm") == "4,7kOhm"
    assert ss.apply_unit_hint("-5", "mV") == "-5mV"
    assert ss.apply_unit_hint("100n", "pF") == "100n"  # own prefix wins
    assert ss.apply_unit_hint("ca. 100", "pF") == "ca. 100"
    assert ss.apply_unit_hint(None, "pF") is None and ss.apply_unit_hint(True, "pF") is True
    assert ss.apply_unit_hint("100", None) == "100" and ss.apply_unit_hint("", "pF") == ""
    assert ss.apply_unit_hint("-40 … 85", "°C", "range") == "-40°C … 85°C"
    assert ss.apply_unit_hint("1.8 to 5.5", "V", "range") == "1.8V to 5.5V"
    assert ss.apply_unit_hint("-40 … 85", "°C", "quantity") == "-40 … 85"


def test_percent_formatted_cells(workspace: Workspace, tmp_path: Path) -> None:
    """Regression: Excel/LibreOffice store "1 %" as 0.01 – it was imported as ±0.01%."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Part Number", "Manufacturer", "Resistance", "Tolerance", "Power", "Library Ref", "Footprint Ref"])
    sheet.append(["R-PCT", "YAGEO", 10000, 0.01, 0.1, "RES", "RESC1608X55N"])
    sheet.append(["R-PCT2", "YAGEO", 4700, 0.005, 0.1, "RES", "RESC1608X55N"])
    sheet.append(["R-LIT", "YAGEO", 1000, 5, 0.1, "RES", "RESC1608X55N"])
    sheet["D2"].number_format = "0%"
    sheet["D3"].number_format = "0.0%"
    sheet["D4"].number_format = '0" %"'  # literal percent sign: the cell value is not scaled
    buffer = io.BytesIO()
    wb.save(buffer)
    token, path = upload(buffer.getvalue(), "pct.xlsx")
    prev = ss.preview(path, table="Resistors")
    assert [r[3] for r in prev["rows"]] == ["1%", "0.5%", "5"]
    result = ss.run_import(None, ImportOptions(token=token, table="Resistors", mapping=prev["mapping"]))
    assert result["created"] == 3, result
    assert parts.get_part("Resistors", "R-PCT")["Tolerance"] == "±1%"  # type: ignore[index]
    assert parts.get_part("Resistors", "R-PCT2")["Tolerance"] == "±0.5%"  # type: ignore[index]
    assert parts.get_part("Resistors", "R-LIT")["Tolerance"] == "±5%"  # type: ignore[index]

    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name="T")
    for values in (("Teilenummer", "Toleranz"), ("R-ODS", 0.01)):
        tr = TableRow()
        for value in values:
            if isinstance(value, float):
                cell = TableCell(valuetype="percentage", value=str(value))
                cell.addElement(P(text="1%"))
            else:
                cell = TableCell(valuetype="string")
                cell.addElement(P(text=value))
            tr.addElement(cell)
        table.addElement(tr)
    doc.spreadsheet.addElement(table)
    ods = tmp_path / "p.ods"
    doc.save(str(ods))
    assert ss.read_table_file(ods).rows == [["R-ODS", "1%"]]


def test_xlsx_export_keeps_formula_like_text(workspace: Workspace) -> None:
    """Regression: "=Value" (Altium string indirection) was written as an Excel formula."""
    import zipfile

    parts.create_part("Resistors", {"Part Number": "R-EQ", "Manufacturer": "YAGEO", "Manufacturer Part Number": "R-EQ",
                                    "Resistance": "10k", "Tolerance": "1", "Power": "0.1", "Library Ref": "RES",
                                    "Footprint Ref": "RESC1608X55N", "Comment": "=Value",
                                    "Notes": "+49 711 123 / -55 bis 155 \x07", "Supplier 1": "@Distributor"})
    cols = ["Part Number", "Comment", "Notes", "Supplier 1", "Operating Temperature"]
    out = ss.render_export("Resistors", fmt="xlsx", columns=cols)
    archive = zipfile.ZipFile(io.BytesIO(out.content))
    xml = archive.read(next(n for n in archive.namelist() if n.startswith("xl/worksheets/"))).decode("utf-8")
    assert "<f>" not in xml
    token, path = upload(out.content, "export.xlsx")
    data = ss.read_table_file(path)
    assert data.rows == [["R-EQ", "=Value", "+49 711 123 / -55 bis 155", "@Distributor", None]]
    prev = ss.preview(path, table="Resistors")
    parts.delete_parts("Resistors", ["R-EQ"], backup=False)
    result = ss.run_import(None, ImportOptions(token=token, table="Resistors", mapping=prev["mapping"], allow_incomplete=True))
    assert result["created"] == 1 and parts.get_part("Resistors", "R-EQ")["Comment"] == "=Value"  # type: ignore[index]


def test_csv_excel_safe_export_roundtrip(workspace: Workspace) -> None:
    parts.create_part("Resistors", {"Part Number": "R-CSV", "Manufacturer": "YAGEO", "Manufacturer Part Number": "R-CSV",
                                    "Resistance": "10k", "Tolerance": "1", "Power": "0.1", "Library Ref": "RES",
                                    "Footprint Ref": "RESC1608X55N", "Comment": "=Value",
                                    "Operating Temperature": "-55 … 155"})
    cols = ["Part Number", "Comment", "Operating Temperature", "Resistance"]
    plain = ss.render_export("Resistors", fmt="csv", columns=cols, excel_safe=False).content.decode("utf-8-sig")
    assert plain.splitlines()[1] == "R-CSV;=Value;-55°C … 155°C;10kΩ"  # excel_safe=False: values unchanged
    safe = ss.render_export("Resistors", fmt="csv", columns=cols).content  # default: guarded (INT-09)
    assert safe.decode("utf-8-sig").splitlines()[1] == "R-CSV;'=Value;-55°C … 155°C;10kΩ"  # a measurement stays
    _, path = upload(safe, "safe.csv")
    assert ss.read_table_file(path).rows == [["R-CSV", "=Value", "-55°C … 155°C", "10kΩ"]]
    assert ss._unguard("'Text") == "'Text" and ss._unguard("'") == "'"


def test_file_without_header_row_like_altium365_export(workspace: Workspace) -> None:
    """Regression: a data row was taken as header (row 1 of an Altium 365 export is data)."""
    rows = [["CMP-428-000063", "CAP CER 68pF 50V ±5% NP0 1210", "PCC-432-000016-1", "SYM-CAP-000002-3"],
            ["CMP-428-000064", "CAP CER 100pF 50V ±5% NP0 1210", "PCC-432-000016-1", "SYM-CAP-000002-3"]]
    _, path = upload("\n".join(";".join(r) for r in rows).encode("utf-8"), "a365.csv")
    data = ss.read_table_file(path, category=categories.get_category("Capacitors"))
    assert data.header_row is None and data.headers[0] == "Spalte A" and len(data.rows) == 2
    # a real header above similar data rows is still recognised
    _, path2 = upload("Teilenummer;Beschreibung\nCMP-428-000063;CAP CER 68pF\nCMP-428-000064;CAP CER 100pF\n".encode(), "mit.csv")
    assert ss.read_table_file(path2).header_row == 1
    assert ss._shape("CMP-428-000063") == "A-9-9" and ss._shape(None) == ""
